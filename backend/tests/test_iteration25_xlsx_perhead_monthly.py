# Iteration 25: xlsx/csv exports, per-head water in cycles, monthly e-mail report (GET/PUT/send, RBAC, events)
import io
import os
import uuid
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo
import pytest
import requests
from openpyxl import load_workbook
from pymongo import MongoClient

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = f"{BASE_URL}/api"
H = {"Origin": BASE_URL}
KOW = "kowalski-a"
KOW_CID = "6ad7dcc3261e5ba3c23315b7"
TAG = uuid.uuid4().hex[:6]
PW = "Haslo1234!x"
NOW = datetime.now(ZoneInfo("Europe/Warsaw")).date()
FIRST = NOW.replace(day=1)
PREV_LAST = FIRST - timedelta(days=1)
PREV_FIRST = PREV_LAST.replace(day=1)
mongo = MongoClient("mongodb://localhost:27017")["test_database"]
# synthetic meter days in previous month (chamber 0 = 1000 l/day) and current month (chamber 0 = 2000 l/day)
PREV_DAYS = [(PREV_FIRST + timedelta(days=i)).isoformat() for i in range(10)]
CUR_DAYS = [(FIRST + timedelta(days=i)).isoformat() for i in range(5) if FIRST + timedelta(days=i) < NOW]


def login(email, pw):
    ss = requests.Session()
    ss.headers.update(H)
    r = ss.post(f"{API}/auth/login", json={"email": email, "password": pw})
    assert r.status_code == 200, r.text
    return ss


@pytest.fixture(scope="module")
def sa():
    s = login("admin@chlewnia.pl", "Wentylacja!2026")
    s.uids, s.emails = [], {}
    for role in ("admin", "operator", "viewer"):
        email = f"test_{TAG}_{role}@kowalski.pl"
        r = s.post(f"{API}/users", json={"email": email, "password": PW, "role": role, "client_id": KOW_CID})
        assert r.status_code in (200, 201), r.text
        s.uids.append(r.json().get("id") or r.json().get("user", {}).get("id"))
        s.emails[role] = email
    inserted = []
    for d in PREV_DAYS + CUR_DAYS:
        _id = f"{KOW}|{d}"
        if not mongo.meter_daily.find_one({"_id": _id}):
            w = 1000 if d in PREV_DAYS else 2000
            mongo.meter_daily.insert_one({"_id": _id, "device_id": KOW, "day": d, "kwh": 10.0,
                                          "water": {str(i): float(w) for i in range(6)}, "test_tag": TAG})
            inserted.append(_id)
    mongo.cycles.delete_many({"device_id": KOW, "chamber": 0})
    yield s
    mongo.meter_daily.delete_many({"_id": {"$in": inserted}})
    mongo.cycles.delete_many({"device_id": KOW, "started_by": {"$regex": f"test_{TAG}"}})
    mongo.cycles.delete_many({"device_id": KOW, "chamber": 0})
    mongo.clients.update_one({"name": "Ferma Kowalski"}, {"$unset": {"monthly_report": ""}})
    mongo.plc_events.delete_many({"message": {"$regex": f"test_{TAG}"}})
    for uid in s.uids:
        s.delete(f"{API}/users/{uid}")


@pytest.fixture(scope="module")
def adm(sa):
    return login(sa.emails["admin"], PW)


@pytest.fixture(scope="module")
def op(sa):
    return login(sa.emails["operator"], PW)


@pytest.fixture(scope="module")
def viewer(sa):
    return login(sa.emails["viewer"], PW)


# --- exports ---
class TestExports:
    def test_history_xlsx(self, sa):
        r = sa.get(f"{API}/devices/{KOW}/history/xlsx", params={"chamber": 0, "hours": 24})
        assert r.status_code == 200, r.text
        assert "spreadsheetml" in r.headers["content-type"]
        assert ".xlsx" in r.headers.get("content-disposition", "")
        ws = load_workbook(io.BytesIO(r.content)).active
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0][0] == "Czas"
        assert len(rows) > 1
        assert ws["A1"].font.bold

    def test_history_xlsx_viewer(self, viewer):
        r = viewer.get(f"{API}/devices/{KOW}/history/xlsx", params={"chamber": 1, "hours": 1})
        assert r.status_code == 200
        load_workbook(io.BytesIO(r.content))

    def test_history_xlsx_invalid_chamber(self, sa):
        assert sa.get(f"{API}/devices/{KOW}/history/xlsx", params={"chamber": 6}).status_code == 404
        assert sa.get(f"{API}/devices/{KOW}/history/xlsx", params={"chamber": -1}).status_code == 404

    def test_history_xlsx_cross_tenant(self, viewer):
        assert viewer.get(f"{API}/devices/chlodnia-01/history/xlsx", params={"chamber": 0}).status_code in (403, 404)

    def test_history_csv(self, viewer):
        r = viewer.get(f"{API}/devices/{KOW}/history/csv", params={"chamber": 0, "hours": 2})
        assert r.status_code == 200
        txt = r.content.decode("utf-8-sig")
        assert txt.startswith("czas;temperatura_C")
        assert len(txt.strip().split("\n")) > 1

    def test_usage_range_csv(self, viewer):
        r = viewer.get(f"{API}/devices/{KOW}/usage/range/csv", params={"start": PREV_FIRST.isoformat(), "end": PREV_LAST.isoformat()})
        assert r.status_code == 200, r.text
        lines = r.content.decode("utf-8-sig").strip().split("\n")
        assert lines[0].startswith("dzien;prad_kWh;woda_")
        assert len(lines[0].split(";")) == 8
        assert lines[-1].startswith("RAZEM;")
        tot = lines[-1].split(";")
        assert tot[2] == "10,0" or tot[2] == "10"  # 10 days * 1000 l = 10 m3
        assert len(lines) == len(PREV_DAYS) + 2

    def test_usage_range_xlsx(self, viewer):
        r = viewer.get(f"{API}/devices/{KOW}/usage/range/xlsx", params={"start": PREV_FIRST.isoformat(), "end": PREV_LAST.isoformat()})
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Podsumowanie", "Dni"]
        assert wb["Podsumowanie"]["B3"].value == 100.0  # 10 days * 10 kWh

    def test_usage_range_csv_invalid(self, viewer):
        assert viewer.get(f"{API}/devices/{KOW}/usage/range/csv", params={"start": "2026-10-10", "end": "2026-10-01"}).status_code == 422
        assert viewer.get(f"{API}/devices/{KOW}/usage/range/csv", params={"start": "bad", "end": "2026-10-01"}).status_code == 422


# --- per-head ---
class TestPerHead:
    def test_per_head(self, sa, op):
        a, b = PREV_DAYS[0], PREV_DAYS[-1]
        mongo.cycles.insert_many([
            {"device_id": KOW, "chamber": 0, "start": a, "end": b, "animals": 100, "started_by": f"test_{TAG}@x"},
            {"device_id": KOW, "chamber": 0, "start": a, "end": PREV_DAYS[4], "animals": None, "started_by": f"test_{TAG}@x"},
        ])
        start = CUR_DAYS[0] if CUR_DAYS else NOW.isoformat()
        r = op.post(f"{API}/devices/{KOW}/chambers/0/cycles", json={"start": start, "animals": 200})
        assert r.status_code == 200, r.text
        act = r.json()
        assert act["water_per_head_l"] == round(act["water_l"] / 200, 2)
        cyc = sa.get(f"{API}/devices/{KOW}/cycles", params={"chamber": 0}).json()["cycles"]
        by = {(c["start"], c["end"]): c for c in cyc}
        ended = by[(a, b)]
        assert ended["water_l"] == 10000
        assert ended["water_per_head_l"] == 100.0
        assert by[(a, PREV_DAYS[4])]["water_per_head_l"] is None
        live = by[(start, None)]
        assert live["animals"] == 200
        assert live["water_per_head_l"] == round(live["water_l"] / 200, 2)
        assert live["water_l"] >= 2000 * len(CUR_DAYS)


# --- monthly report ---
class TestMonthly:
    def test_get(self, sa):
        r = sa.get(f"{API}/monthly-report", params={"client_id": KOW_CID})
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["enabled"] is True
        assert sa.emails["admin"] in d["recipients"]
        assert sa.emails["operator"] not in d["recipients"]
        assert d["period"] == [PREV_FIRST.isoformat(), PREV_LAST.isoformat()]

    def test_superadmin_requires_client(self, sa):
        assert sa.get(f"{API}/monthly-report").status_code == 422

    def test_put_persist(self, sa):
        r = sa.put(f"{API}/monthly-report", params={"client_id": KOW_CID}, json={"enabled": False})
        assert r.status_code == 200 and r.json()["enabled"] is False
        assert sa.get(f"{API}/monthly-report", params={"client_id": KOW_CID}).json()["enabled"] is False
        assert mongo.clients.find_one({"name": "Ferma Kowalski"})["monthly_report"] is False
        sa.put(f"{API}/monthly-report", params={"client_id": KOW_CID}, json={"enabled": True})
        assert sa.get(f"{API}/monthly-report", params={"client_id": KOW_CID}).json()["enabled"] is True

    def test_put_invalid(self, sa):
        assert sa.put(f"{API}/monthly-report", params={"client_id": KOW_CID}, json={"enabled": "maybe"}).status_code == 422

    def test_send(self, sa):
        r = sa.post(f"{API}/monthly-report/send", params={"client_id": KOW_CID})
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["sent"] == 0
        assert sa.emails["admin"] in d["recipients"]
        assert d["status"] == ["not_configured"] * len(d["recipients"])
        ev = sa.get(f"{API}/events", params={"limit": 50}).json()["events"]
        msg = f"Raport miesięczny {PREV_FIRST.isoformat()[:7]}:"
        hit = [e for e in ev if e["message"].startswith(msg) and sa.emails["admin"] in e["message"]]
        assert hit, [e["message"] for e in ev[:5]]
        assert hit[0]["client_id"] == KOW_CID
        rep = mongo.monthly_reports.find_one({"_id": f"{KOW_CID}|{PREV_FIRST.isoformat()[:7]}"})
        assert rep and rep["status"] == d["status"]

    def test_client_admin_own_scope(self, adm):
        r = adm.get(f"{API}/monthly-report")
        assert r.status_code == 200
        assert r.json()["enabled"] is True
        r = adm.post(f"{API}/monthly-report/send")
        assert r.status_code == 200 and r.json()["status"][0] == "not_configured"

    def test_admin_ignores_foreign_client(self, adm, sa):
        mp = sa.get(f"{API}/monthly-report", params={"client_id": "6ad7dcb08b38a7c48f3db7d2"}).json()
        own = adm.get(f"{API}/monthly-report", params={"client_id": "6ad7dcb08b38a7c48f3db7d2"}).json()
        assert sa.emails["admin"] in own["recipients"]
        r = adm.put(f"{API}/monthly-report", params={"client_id": "6ad7dcb08b38a7c48f3db7d2"}, json={"enabled": True})
        assert r.status_code == 200
        assert mongo.clients.find_one({"name": "Marco Polo"}).get("monthly_report", True) == mp["enabled"]

    @pytest.mark.parametrize("role", ["operator", "viewer"])
    def test_rbac(self, role, op, viewer):
        s = op if role == "operator" else viewer
        assert s.get(f"{API}/monthly-report").status_code == 403
        assert s.put(f"{API}/monthly-report", json={"enabled": False}).status_code == 403
        assert s.post(f"{API}/monthly-report/send").status_code == 403

    def test_unauth(self):
        assert requests.get(f"{API}/monthly-report", params={"client_id": KOW_CID}).status_code == 401
