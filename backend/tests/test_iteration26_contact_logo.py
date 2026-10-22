# Iteration 26: public contact form (validation/honeypot/rate-limit/503), site settings (address/website), xlsx logo embedding, monthly report send
import io
import os
import pytest
import requests
from openpyxl import load_workbook
from pymongo import MongoClient

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = f"{BASE_URL}/api"
H = {"Origin": BASE_URL}
KOW = "kowalski-a"
mongo = MongoClient("mongodb://localhost:27017")["test_database"]
IP = "203.0.113.77"


def login(email, pw):
    ss = requests.Session()
    ss.headers.update(H)
    r = ss.post(f"{API}/auth/login", json={"email": email, "password": pw})
    assert r.status_code == 200, r.text
    tok = ss.cookies.get("csrf_token") or ss.cookies.get("csrf")
    if tok:
        ss.headers["X-CSRF-Token"] = tok
    return ss


@pytest.fixture(scope="module")
def sa():
    s = login("admin@chlewnia.pl", "Wentylacja!2026")
    orig = requests.get(f"{API}/public/site").json()
    yield s
    s.put(f"{API}/site", json={k: orig[k] for k in ("phone", "email", "address", "website")})
    mongo.contact_requests.delete_many({})


def post_contact(body, ip=IP):
    return requests.post(f"{API}/public/contact", json=body, headers={**H, "X-Forwarded-For": ip})


GOOD = {"name": "TEST Jan", "phone": "+48 600 000 000", "email": "", "message": "Prosze o kontakt", "company": ""}


# --- contact form ---
class TestContact:
    def test_short_name_422(self, sa):
        mongo.contact_requests.delete_many({})
        assert post_contact({**GOOD, "name": "J"}).status_code == 422

    def test_short_message_422(self, sa):
        assert post_contact({**GOOD, "message": "abc"}).status_code == 422

    def test_no_phone_no_email_422(self, sa):
        assert post_contact({**GOOD, "phone": "", "email": ""}).status_code == 422

    def test_bad_email_422(self, sa):
        assert post_contact({**GOOD, "email": "nope"}).status_code == 422

    def test_honeypot_200(self, sa):
        r = post_contact({**GOOD, "company": "bot inc"})
        assert r.status_code == 200 and r.json() == {"ok": True}
        assert mongo.contact_requests.count_documents({"ip": IP}) == 0

    def test_real_send_503_then_rate_limit(self, sa):
        mongo.contact_requests.delete_many({})
        codes = [post_contact(GOOD).status_code for _ in range(4)]
        assert codes[:3] == [503, 503, 503], codes
        assert codes[3] == 429, codes
        # different IP not limited
        assert post_contact(GOOD, ip="203.0.113.78").status_code == 503
        mongo.contact_requests.delete_many({})

    def test_no_site_email_503(self, sa):
        orig = requests.get(f"{API}/public/site").json()
        r = sa.put(f"{API}/site", json={**{k: orig[k] for k in ("phone", "address", "website")}, "email": ""})
        assert r.status_code == 200
        assert post_contact(GOOD, ip="203.0.113.79").status_code == 503
        sa.put(f"{API}/site", json={k: orig[k] for k in ("phone", "email", "address", "website")})


# --- site settings ---
class TestSite:
    def test_put_site_persists(self, sa):
        body = {"phone": "+48 111 222 333", "email": "test_biuro@example.pl", "address": "ul. Testowa 5, 11-111 Test",
                "website": "https://test.example.pl"}
        r = sa.put(f"{API}/site", json=body)
        assert r.status_code == 200, r.text
        g = requests.get(f"{API}/public/site").json()
        for k, v in body.items():
            assert g[k] == v

    def test_put_site_invalid(self, sa):
        assert sa.put(f"{API}/site", json={"phone": "abc", "email": "", "address": "", "website": ""}).status_code == 422
        assert sa.put(f"{API}/site", json={"phone": "", "email": "", "address": "", "website": "a b<script>"}).status_code == 422

    def test_put_site_unauth(self):
        assert requests.put(f"{API}/site", json={"phone": ""}, headers=H).status_code in (401, 403)


def check_logo_xlsx(content):
    wb = load_workbook(io.BytesIO(content))
    for ws in wb.worksheets:
        assert len(ws._images) == 1, ws.title
        assert ws._images[0].anchor._from.row == 0
        for r in range(1, 6):
            assert all(c.value is None for c in ws[r])
        assert ws.cell(6, 1).value is not None and ws.cell(6, 1).font.bold
    return wb


# --- xlsx exports with logo ---
class TestXlsx:
    def test_cycles_export(self, sa):
        r = sa.get(f"{API}/devices/{KOW}/cycles/export")
        assert r.status_code == 200
        wb = check_logo_xlsx(r.content)
        assert wb.worksheets[0].cell(6, 1).value == "Komora"

    def test_usage_range(self, sa):
        r = sa.get(f"{API}/devices/{KOW}/usage/range/xlsx", params={"start": "2026-06-01", "end": "2026-06-10"})
        assert r.status_code == 200, r.text
        check_logo_xlsx(r.content)

    def test_history(self, sa):
        r = sa.get(f"{API}/devices/{KOW}/history/xlsx", params={"chamber": 0, "hours": 2})
        assert r.status_code == 200
        wb = check_logo_xlsx(r.content)
        assert wb.worksheets[0].cell(6, 1).value == "Czas"

    def test_csv_unchanged(self, sa):
        r = sa.get(f"{API}/devices/{KOW}/cycles/export", params={"fmt": "csv"})
        assert r.status_code == 200 and r.content.decode("utf-8-sig").startswith("Komora;")
        r = sa.get(f"{API}/devices/{KOW}/usage/range/csv", params={"start": "2026-06-01", "end": "2026-06-10"})
        assert r.status_code == 200 and r.content.decode("utf-8-sig").startswith("dzien;")


# --- monthly report ---
def test_monthly_send(sa):
    c = mongo.clients.find_one({"name": "Ferma Kowalski"})
    r = sa.post(f"{API}/monthly-report/send", params={"client_id": str(c["_id"])})
    assert r.status_code == 200, r.text
    assert "recipients" in r.json() or "sent" in r.json()


# --- fallback Marco-Polo logo for client without logo (chlodnia-01) vs client logo (kowalski-a) ---
def test_logo_fallback_vs_client(sa):
    a = sa.get(f"{API}/devices/chlodnia-01/history/xlsx", params={"chamber": 0, "hours": 1})
    b = sa.get(f"{API}/devices/{KOW}/history/xlsx", params={"chamber": 0, "hours": 1})
    assert a.status_code == 200 and b.status_code == 200
    ia = load_workbook(io.BytesIO(a.content)).worksheets[0]._images
    ib = load_workbook(io.BytesIO(b.content)).worksheets[0]._images
    assert len(ia) == 1 and len(ib) == 1
    assert ia[0]._data() != ib[0]._data()
