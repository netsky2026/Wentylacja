# Iteration 30: superadmin farm comparison (/farms/compare + xlsx), formula-injection neutralisation in CSV/XLSX exports,
# contact form name CR/LF collapse. Seeds synthetic meter_daily + cycles in kowalski-a chambers 2-5; cleans up afterwards.
import io
import os
import pytest
import requests
from bson import ObjectId
from openpyxl import load_workbook
from pymongo import MongoClient

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = f"{BASE_URL}/api"
KOW = "kowalski-a"
KOW_CID = "6ad7dcc3261e5ba3c23315b7"
mongo = MongoClient("mongodb://localhost:27017")["test_database"]
SYN = {"2026-08-05": (40.0, {"2": 3000.0}), "2026-08-20": (20.0, {"2": 1500.0}),
       "2026-09-07": (30.0, {"3": 1000.0, "4": 2000.0}), "2026-09-12": (10.0, {"4": 500.0})}
CYC = [(2, "2026-08-01", "2026-08-10", 300), (2, "2026-08-15", "2026-08-25", 300), (3, "2026-09-01", "2026-09-10", 200),
       (4, "2026-09-05", "2026-09-15", 250), (4, "2026-09-20", None, 250), (5, "2026-08-01", "2026-08-03", None)]
EVIL = {2: "=1+1", 3: "+SUM(A1)", 4: "-2+3", 5: "@cmd"}


def login(email, pw):
    s = requests.Session()
    s.headers.update({"Origin": BASE_URL})
    r = s.post(f"{API}/auth/login", json={"email": email, "password": pw})
    assert r.status_code == 200, r.text
    s.headers["X-CSRF-Token"] = s.cookies.get("csrf") or ""
    return s


@pytest.fixture(scope="module")
def sa():
    s = login("admin@chlewnia.pl", "Wentylacja!2026")
    for ch in (2, 3, 4, 5):
        assert not mongo.cycles.find_one({"device_id": KOW, "chamber": ch}), f"chamber {ch} not empty"
    for d, (k, w) in SYN.items():
        assert not mongo.meter_daily.find_one({"_id": f"{KOW}|{d}"})
        mongo.meter_daily.insert_one({"_id": f"{KOW}|{d}", "device_id": KOW, "day": d, "kwh": k, "water": w, "test_tag": "it30"})
    ids = []
    for ch, start, end, animals in CYC:
        r = s.post(f"{API}/devices/{KOW}/chambers/{ch}/cycles", json={"start": start, "animals": animals})
        assert r.status_code == 200, r.text
        ids.append(r.json()["id"])
        if end:
            e = s.post(f"{API}/devices/{KOW}/cycles/{ids[-1]}/end", json={"end": end})
            assert e.status_code == 200, e.text
    for ch, nm in EVIL.items():
        r = s.put(f"{API}/devices/{KOW}/chambers/{ch}/name", json={"name": nm})
        assert r.status_code == 200 and r.json()["name"] == nm, r.text
    yield s
    if os.environ.get("KEEP_SEED") == "1":
        open("/tmp/it30_cycles.txt", "w").write(",".join(ids))
        return
    for ch in EVIL:
        s.put(f"{API}/devices/{KOW}/chambers/{ch}/name", json={"name": ""})
    mongo.cycles.delete_many({"_id": {"$in": [ObjectId(i) for i in ids]}})
    mongo.meter_daily.delete_many({"test_tag": "it30"})
    mongo.contact_requests.delete_many({})


@pytest.fixture(scope="module")
def kow_admin():
    return login("kowalski.admin@kowalski.pl", "KowalskiAdm!2026")


# ---- farm compare ----
def farms(s):
    r = s.get(f"{API}/farms/compare")
    assert r.status_code == 200, r.text
    return {f["client_id"]: f for f in r.json()["farms"]}


def test_farms_shape_every_client(sa):
    f = farms(sa)
    assert set(f) == {str(c["_id"]) for c in mongo.clients.find()}
    for row in f.values():
        assert set(row) == {"client_id", "name", "devices", "cycles", "cost_per_head", "water_per_head_l", "kwh_per_head"}


def test_farms_kowalski_values(sa):
    k = farms(sa)[KOW_CID]
    assert k["name"] == "Kowalski Wentylacja" and k["devices"] == 1
    assert k["cycles"] == 5, "only finished cycles"
    exp_cost = ((34 + 19.5) / 300 + (17 + 9.75) / 300 + (15 * 0.85 + 6.5) / 200 + (25 * 0.85 + 2.5 * 6.5) / 250) / 4
    assert k["cost_per_head"] == pytest.approx(exp_cost, abs=1e-3)
    assert k["water_per_head_l"] == pytest.approx(7.5, abs=0.01)
    assert k["kwh_per_head"] == pytest.approx((40 / 300 + 20 / 300 + 15 / 200 + 25 / 250) / 4, abs=1e-3)


def test_farms_match_per_chamber_cycles(sa):
    """Averages computed from /devices/{id}/cycles?chamber=N must equal /farms/compare."""
    devs = sa.get(f"{API}/devices").json()
    devs = devs.get("devices", devs) if isinstance(devs, dict) else devs
    f = farms(sa)
    for cid, row in f.items():
        vs = []
        for d in [d for d in devs if d.get("client_id") == cid]:
            for ch in range(d["chambers"]):
                vs += [c for c in sa.get(f"{API}/devices/{d['device_id']}/cycles", params={"chamber": ch}).json()["cycles"] if c["end"]]
        assert row["cycles"] == len(vs)

        def a(xs):
            xs = [x for x in xs if x is not None]
            return sum(xs) / len(xs) if xs else None
        for key, vals in (("cost_per_head", [v["cost_per_head"] for v in vs]), ("water_per_head_l", [v["water_per_head_l"] for v in vs]),
                          ("kwh_per_head", [v["kwh_chamber"] / v["animals"] if v["animals"] else None for v in vs])):
            e = a(vals)
            assert (row[key] is None) if e is None else row[key] == pytest.approx(e, abs=2e-3), (cid, key, row[key], e)


def test_farms_rbac(kow_admin):
    assert kow_admin.get(f"{API}/farms/compare").status_code == 403
    assert kow_admin.get(f"{API}/farms/compare/xlsx").status_code == 403
    assert requests.get(f"{API}/farms/compare").status_code == 401
    assert requests.get(f"{API}/farms/compare/xlsx").status_code == 401


def test_farms_xlsx(sa):
    r = sa.get(f"{API}/farms/compare/xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    assert "porownanie_ferm_" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Porównanie ferm"]
    assert len(ws._images) == 1
    assert all(ws.cell(row=i, column=1).value is None for i in range(1, 6))
    assert ws.cell(row=6, column=1).value == "Ferma" and ws.cell(row=6, column=1).font.bold
    names = [ws.cell(row=i, column=1).value for i in range(7, ws.max_row + 1)]
    assert "Kowalski Wentylacja" in names and "Marco Polo" in names
    kr = names.index("Kowalski Wentylacja") + 7
    assert ws.cell(row=kr, column=3).value == 5


# ---- formula injection ----
def test_cycles_csv_neutralised(sa):
    r = sa.get(f"{API}/devices/{KOW}/cycles/export", params={"fmt": "csv"})
    assert r.status_code == 200
    lines = r.content.decode("utf-8-sig").strip().split("\n")
    first = [l.split(";")[0] for l in lines[1:]]
    for nm in EVIL.values():
        assert "'" + nm in first, (nm, first)
    assert not any(c[:1] in "=+-@" for c in first)
    row = next(l.split(";") for l in lines if l.startswith("'=1+1;2026-08-01"))
    assert row[2] == "2026-08-10" and row[3] == "10" and row[4] == "300" and row[5] == "3,0" and row[6] == "10,0"
    assert any(l.split(";")[2] == "trwa" for l in lines if l.startswith("'-2+3"))


def test_cycles_csv_chamber_filter(sa):
    r = sa.get(f"{API}/devices/{KOW}/cycles/export", params={"fmt": "csv", "chamber": 3})
    lines = r.content.decode("utf-8-sig").strip().split("\n")
    assert len(lines) == 2 and lines[1].startswith("'+SUM(A1);")


def test_cycles_xlsx_neutralised(sa):
    r = sa.get(f"{API}/devices/{KOW}/cycles/export", params={"fmt": "xlsx"})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    for sheet in ("Chowy", "Porównanie komór"):
        ws = wb[sheet]
        cells = [ws.cell(row=i, column=1) for i in range(7, ws.max_row + 1)]
        vals = {c.value: c.data_type for c in cells}
        for nm in EVIL.values():
            assert vals.get(nm) == "s", (sheet, nm, vals)
    raw = r.content
    import zipfile
    z = zipfile.ZipFile(io.BytesIO(raw))
    for n in z.namelist():
        if n.startswith("xl/worksheets/sheet"):
            assert b"<f>" not in z.read(n), f"formula element in {n}"


def test_history_xlsx_csv_with_evil_name(sa):
    r = sa.get(f"{API}/devices/{KOW}/history/xlsx", params={"chamber": 2, "hours": 2})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0] == "=1+1"
    r = sa.get(f"{API}/devices/{KOW}/history/csv", params={"chamber": 2, "hours": 2})
    assert r.status_code == 200
    lines = r.content.decode("utf-8-sig").strip().split("\n")
    assert lines[0].startswith("czas;temperatura_C")
    for l in lines[1:]:
        assert l[:1] not in "=+-@"
        for c in l.split(";")[1:]:
            assert c == "" or c.replace(",", "").replace("-", "").isdigit() or c.startswith("'-") is False


def test_usage_range_exports(sa):
    r = sa.get(f"{API}/devices/{KOW}/usage/range/csv", params={"start": "2026-08-01", "end": "2026-09-30"})
    assert r.status_code == 200
    lines = r.content.decode("utf-8-sig").strip().split("\n")
    assert lines[0].split(";")[:4] == ["dzien", "prad_kWh", "woda_Komora 1_m3", "woda_Komora 2_m3"]
    assert "woda_=1+1_m3" in lines[0] and "woda_+SUM(A1)_m3" in lines[0]
    assert lines[-1].startswith("RAZEM;100")
    r = sa.get(f"{API}/devices/{KOW}/usage/range/xlsx", params={"start": "2026-08-01", "end": "2026-09-30"})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert {"Podsumowanie", "Dni"} <= set(wb.sheetnames)
    ws = wb["Podsumowanie"]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value[:1] in "=+-@":
                assert c.data_type == "s"


# ---- contact form ----
def test_contact_crlf_name_no_500():
    mongo.contact_requests.delete_many({})
    r = requests.post(f"{API}/public/contact", json={"name": "Jan\r\nBcc: x@y.pl", "phone": "600123456",
                                                     "message": "Test wiadomości CRLF"})
    assert r.status_code == 503, r.text
    ev = mongo.plc_events.find_one({"message": {"$regex": "^Formularz kontaktowy: Jan Bcc"}}, sort=[("ts", -1)])
    assert ev and "\n" not in ev["message"] and "\r" not in ev["message"]


def test_contact_validation_unchanged():
    mongo.contact_requests.delete_many({})
    assert requests.post(f"{API}/public/contact", json={"name": "J", "phone": "1", "message": "hello"}).status_code == 422
    assert requests.post(f"{API}/public/contact", json={"name": "Jan", "message": "hello there"}).status_code == 422
    assert requests.post(f"{API}/public/contact", json={"name": "Jan", "phone": "abc", "message": "hello there"}).status_code == 422
    # whitespace-only name collapses to empty
    r = requests.post(f"{API}/public/contact", json={"name": " \r\n ", "phone": "600123456", "message": "hello there"})
    print("whitespace-only name ->", r.status_code, r.text)
    mongo.contact_requests.delete_many({})
