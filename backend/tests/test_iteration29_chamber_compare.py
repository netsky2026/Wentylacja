# Iteration 29: chamber comparison (GET /cycles/compare) averaged over finished cycles + xlsx 'Porównanie komór' sheet.
# Seeds synthetic meter_daily + cycles in chambers 2-5 of kowalski-a; cleanup unless KEEP_SEED=1.
import io
import json
import os
import pytest
import requests
from bson import ObjectId
from pymongo import MongoClient

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = f"{BASE_URL}/api"
KOW = "kowalski-a"
mongo = MongoClient("mongodb://localhost:27017")["test_database"]
SYN = {"2026-08-05": (40.0, {"2": 3000.0}), "2026-08-20": (20.0, {"2": 1500.0}),
       "2026-09-07": (30.0, {"3": 1000.0, "4": 2000.0}), "2026-09-12": (10.0, {"4": 500.0})}
# (chamber, start, end, animals)
CYC = [(2, "2026-08-01", "2026-08-10", 300), (2, "2026-08-15", "2026-08-25", 300), (3, "2026-09-01", "2026-09-10", 200),
       (4, "2026-09-05", "2026-09-15", 250), (4, "2026-09-20", None, 250), (5, "2026-08-01", "2026-08-03", None)]


def login(email, pw):
    s = requests.Session()
    s.headers.update({"Origin": BASE_URL})
    r = s.post(f"{API}/auth/login", json={"email": email, "password": pw})
    assert r.status_code == 200, r.text
    s.headers["X-CSRF-Token"] = s.cookies.get("csrf") or ""
    return s


@pytest.fixture(scope="module")
def sa():
    s = login("admin@chlewnia.pl", "Wentylacja!2026")
    for ch in (2, 3, 4, 5):
        assert not mongo.cycles.find_one({"device_id": KOW, "chamber": ch}), f"chamber {ch} not empty"
    for d, (k, w) in SYN.items():
        assert not mongo.meter_daily.find_one({"_id": f"{KOW}|{d}"})
        mongo.meter_daily.insert_one({"_id": f"{KOW}|{d}", "device_id": KOW, "day": d, "kwh": k, "water": w, "test_tag": "it29"})
    ids = []
    for ch, start, end, animals in CYC:
        r = s.post(f"{API}/devices/{KOW}/chambers/{ch}/cycles", json={"start": start, "animals": animals})
        assert r.status_code == 200, r.text
        ids.append(r.json()["id"])
        if end:
            e = s.post(f"{API}/devices/{KOW}/cycles/{ids[-1]}/end", json={"end": end})
            assert e.status_code == 200, e.text
    json.dump(ids, open("/tmp/it29_cycles.json", "w"))
    yield s
    if os.environ.get("KEEP_SEED") != "1":
        mongo.cycles.delete_many({"_id": {"$in": [ObjectId(i) for i in ids]}})
        mongo.meter_daily.delete_many({"test_tag": "it29"})


def by_ch(s):
    r = s.get(f"{API}/devices/{KOW}/cycles/compare")
    assert r.status_code == 200, r.text
    return {c["chamber"]: c for c in r.json()["chambers"]}


def test_compare_shape_all_chambers(sa):
    rows = by_ch(sa)
    assert sorted(rows) == [0, 1, 2, 3, 4, 5]
    for c in rows.values():
        assert set(c) == {"chamber", "name", "cycles", "cost_per_head", "water_per_head_l", "days"}
        assert isinstance(c["name"], str) and c["name"]


def test_compare_values(sa):
    rows = by_ch(sa)
    c2, c3, c4, c5 = rows[2], rows[3], rows[4], rows[5]
    assert c2["cycles"] == 2
    assert c2["cost_per_head"] == pytest.approx(((34 + 19.5) / 300 + (17 + 9.75) / 300) / 2, abs=1e-3)
    assert c2["water_per_head_l"] == pytest.approx(7.5) and c2["days"] == 10.5
    # ch3 & ch4 overlap on 09-07 -> kwh split in half
    assert c3["cycles"] == 1 and c3["cost_per_head"] == pytest.approx((15 * 0.85 + 6.5) / 200, abs=1e-3)
    assert c4["cycles"] == 1, "active cycle must not be counted"
    assert c4["cost_per_head"] == pytest.approx((25 * 0.85 + 2.5 * 6.5) / 250, abs=1e-3)
    assert c4["water_per_head_l"] == pytest.approx(10) and c4["days"] == 11
    assert c5["cycles"] == 1 and c5["cost_per_head"] is None and c5["water_per_head_l"] is None


def test_compare_empty_chambers_null(sa):
    rows = by_ch(sa)
    for ch in (0, 1):  # only active cycles
        assert rows[ch]["cycles"] == 0
        assert rows[ch]["cost_per_head"] is None and rows[ch]["water_per_head_l"] is None and rows[ch]["days"] is None


def test_compare_matches_list_endpoint(sa):
    rows = by_ch(sa)
    for ch in (2, 3, 4):
        cy = sa.get(f"{API}/devices/{KOW}/cycles?chamber={ch}").json()["cycles"]
        fin = [c["cost_per_head"] for c in cy if c["end"] and c["cost_per_head"] is not None]
        assert rows[ch]["cost_per_head"] == pytest.approx(sum(fin) / len(fin), abs=1e-3)


def test_compare_tenant_scope(sa):
    assert sa.get(f"{API}/devices/nope-xyz/cycles/compare").status_code in (403, 404)
    u = login("kowalski.admin@kowalski.pl", "KowalskiAdm!2026")
    assert u.get(f"{API}/devices/{KOW}/cycles/compare").status_code == 200
    assert u.get(f"{API}/devices/chlodnia-01/cycles/compare").status_code in (403, 404)
    assert requests.get(f"{API}/devices/{KOW}/cycles/compare").status_code == 401


def test_xlsx_two_sheets(sa):
    from openpyxl import load_workbook
    r = sa.get(f"{API}/devices/{KOW}/cycles/export?fmt=xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Chowy", "Porównanie komór"]
    ws = wb["Porównanie komór"]
    head = [c.value for c in ws[6]][:5]
    assert head == ["Komora", "Zakończone rzuty", "Śr. koszt na sztukę zł", "Śr. woda na sztukę l", "Śr. długość rzutu dni"]
    data = [[c.value for c in row][:5] for row in ws.iter_rows(min_row=7)]
    assert len(data) == 6
    assert data[2][1] == 2 and data[2][2] == pytest.approx(0.134, abs=1e-3) and data[0][1] == 0 and data[0][2] is None
    assert any(img for img in getattr(ws, "_images", [])) or ws.max_row >= 7


def test_csv_unchanged(sa):
    r = sa.get(f"{API}/devices/{KOW}/cycles/export?fmt=csv")
    assert r.status_code == 200 and "text/csv" in r.headers["content-type"]
    lines = r.content.decode("utf-8-sig").strip().split("\n")
    assert lines[0].startswith("Komora;Wstawienie;Koniec") and "Porównanie" not in r.text
    assert len(lines) == 1 + mongo.cycles.count_documents({"device_id": KOW})
