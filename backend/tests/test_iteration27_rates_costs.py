# Iteration 27: client rates (kWh, m3), per-chamber kWh allocation, cycle costs, exports with costs, monthly report
import io
import os
import uuid
import pytest
import requests
from openpyxl import load_workbook
from pymongo import MongoClient

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = f"{BASE_URL}/api"
H = {"Origin": BASE_URL}
KOW = "kowalski-a"
CID = "6ad7dcc3261e5ba3c23315b7"
TAG = uuid.uuid4().hex[:6]
PW = "Haslo1234!x"
mongo = MongoClient("mongodb://localhost:27017")["test_database"]
# synthetic days: 10-19, 10-20 only chamber 0 active (start 10-19), chamber 1 starts 10-21
SYN = {"2026-10-19": 20.0, "2026-10-20": 30.0, "2026-10-15": 50.0}


def login(email, pw):
    s = requests.Session()
    s.headers.update(H)
    r = s.post(f"{API}/auth/login", json={"email": email, "password": pw})
    assert r.status_code == 200, r.text
    s.headers["X-CSRF-Token"] = s.cookies.get("csrf") or ""
    return s


@pytest.fixture(scope="module")
def sa():
    s = login("admin@chlewnia.pl", "Wentylacja!2026")
    s.users = {}
    ids = []
    for role in ("admin", "operator", "viewer"):
        email = f"test_{TAG}_{role}@kowalski.pl"
        r = s.post(f"{API}/users", json={"email": email, "password": PW, "role": role, "client_id": CID})
        assert r.status_code in (200, 201), r.text
        ids.append(r.json().get("id") or r.json().get("user", {}).get("id"))
        s.users[role] = email
    inserted = []
    for d, k in SYN.items():
        _id = f"{KOW}|{d}"
        if not mongo.meter_daily.find_one({"_id": _id}):
            mongo.meter_daily.insert_one({"_id": _id, "device_id": KOW, "day": d, "kwh": k,
                                          "water": {str(i): 1000.0 for i in range(6)}, "test_tag": TAG})
            inserted.append(_id)
    yield s
    mongo.meter_daily.delete_many({"_id": {"$in": inserted}})
    s.put(f"{API}/rates?client_id={CID}", json={"kwh": 0.85, "m3": 6.5})
    for i in ids:
        if i:
            s.delete(f"{API}/users/{i}")


# --- rates RBAC + validation
def test_rates_rbac(sa):
    for role, ok in (("admin", True), ("operator", False), ("viewer", False)):
        u = login(sa.users[role], PW)
        assert (u.get(f"{API}/rates").status_code == 200) == ok, role
        r = u.put(f"{API}/rates", json={"kwh": 0.9, "m3": 7})
        assert (r.status_code == 200) == ok, (role, r.status_code)
        if not ok:
            assert r.status_code == 403
    # admin set value visible to superadmin
    assert sa.get(f"{API}/rates?client_id={CID}").json() == {"kwh": 0.9, "m3": 7.0}


def test_rates_validation(sa):
    for body in ({"kwh": -1, "m3": 1}, {"kwh": 1, "m3": -0.1}, {"kwh": 101, "m3": 1}, {"kwh": 1, "m3": 1001}):
        assert sa.put(f"{API}/rates?client_id={CID}", json=body).status_code == 422, body
    r = sa.put(f"{API}/rates?client_id={CID}", json={"kwh": None, "m3": None})
    assert r.status_code == 200 and r.json() == {"kwh": None, "m3": None}
    assert sa.get(f"{API}/rates?client_id={CID}").json() == {"kwh": None, "m3": None}


def test_null_rates_cost_fields_null(sa):
    sa.put(f"{API}/rates?client_id={CID}", json={"kwh": None, "m3": None})
    for ch in (0, 1):
        cy = sa.get(f"{API}/devices/{KOW}/cycles?chamber={ch}").json()["cycles"][0]
        assert cy["kwh_chamber"] is not None
        for k in ("cost_water", "cost_power", "cost_total", "cost_per_head"):
            assert cy[k] is None, (ch, k)


def test_usage_xlsx_no_costs_when_null(sa):
    sa.put(f"{API}/rates?client_id={CID}", json={"kwh": None, "m3": None})
    r = sa.get(f"{API}/devices/{KOW}/usage/range/xlsx?start=2026-06-01&end=2026-10-22")
    assert r.status_code == 200
    vals = [c for row in load_workbook(io.BytesIO(r.content))["Podsumowanie"].iter_rows(values_only=True) for c in row]
    assert "Koszty" not in vals and "KOSZT RAZEM" not in vals


def expected_alloc():
    days = {d["day"]: d.get("kwh", 0) for d in mongo.meter_daily.find({"device_id": KOW})}
    cyc = list(mongo.cycles.find({"device_id": KOW}))
    out = {}
    for c in cyc:
        t = 0.0
        for day, k in days.items():
            if c["start"] <= day <= (c.get("end") or "2026-10-22") and k:
                n = sum(1 for x in cyc if x["start"] <= day <= (x.get("end") or "2026-10-22"))
                t += k / n
        out[c["chamber"]] = round(t, 2)
    return out


def test_cycle_costs_and_allocation(sa):
    sa.put(f"{API}/rates?client_id={CID}", json={"kwh": 0.85, "m3": 6.5})
    c0 = sa.get(f"{API}/devices/{KOW}/cycles?chamber=0").json()["cycles"][0]
    c1 = sa.get(f"{API}/devices/{KOW}/cycles?chamber=1").json()["cycles"][0]
    exp = expected_alloc()
    assert c0["start"] == "2026-10-19" and c0["animals"] == 500
    assert c1["start"] == "2026-10-21" and c1["animals"] == 400
    # synthetic: ch0 gets 20+30 alone; 10-15 is before both cycles -> excluded
    assert abs(c0["kwh_chamber"] - exp[0]) < 0.2, (c0, exp)
    assert abs(c1["kwh_chamber"] - exp[1]) < 0.2, (c1, exp)
    assert abs((c0["kwh_chamber"] - c1["kwh_chamber"]) - 50.0) < 0.05
    for c in (c0, c1):
        assert c["cost_water"] == round(c["water_l"] / 1000 * 6.5, 2)
        assert abs(c["cost_power"] - c["kwh_chamber"] * 0.85) <= 0.011
        assert abs(c["cost_total"] - (c["cost_water"] + c["cost_power"])) <= 0.011
        assert abs(c["cost_per_head"] - c["cost_total"] / c["animals"]) <= 0.011


def test_only_kwh_rate(sa):
    sa.put(f"{API}/rates?client_id={CID}", json={"kwh": 1.0, "m3": None})
    c = sa.get(f"{API}/devices/{KOW}/cycles?chamber=0").json()["cycles"][0]
    assert c["cost_water"] is None and c["cost_power"] is not None
    assert abs(c["cost_total"] - c["cost_power"]) < 0.001
    sa.put(f"{API}/rates?client_id={CID}", json={"kwh": 0.85, "m3": 6.5})


NEW_COLS = ["Prąd kWh (komora)", "Koszt wody zł", "Koszt prądu zł", "Koszt razem zł", "Koszt na sztukę zł"]


def test_cycle_export_xlsx_csv(sa):
    sa.put(f"{API}/rates?client_id={CID}", json={"kwh": 0.85, "m3": 6.5})
    r = sa.get(f"{API}/devices/{KOW}/cycles/export")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content))["Chowy"]
    head = [c.value for c in ws[6]]
    for col in NEW_COLS:
        assert col in head, col
    assert ws.cell(6, 1).font.bold
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=7)]
    assert len(rows) >= 2
    i = head.index("Koszt razem zł")
    assert all(isinstance(rw[i], (int, float)) for rw in rows)
    c = sa.get(f"{API}/devices/{KOW}/cycles/export?fmt=csv")
    assert c.status_code == 200
    first = c.content.decode("utf-8-sig").splitlines()[0].split(";")
    for col in NEW_COLS:
        assert col in first, col


def test_usage_xlsx_costs(sa):
    sa.put(f"{API}/rates?client_id={CID}", json={"kwh": 0.85, "m3": 6.5})
    r = sa.get(f"{API}/devices/{KOW}/usage/range/xlsx?start=2026-06-01&end=2026-10-22")
    assert r.status_code == 200
    rows = list(load_workbook(io.BytesIO(r.content))["Podsumowanie"].iter_rows(values_only=True))
    labels = {rw[0]: rw for rw in rows if rw and rw[0]}
    for k in ("Koszty", "Koszt prądu", "Koszt wody razem", "KOSZT RAZEM"):
        assert k in labels, k
    assert sum(1 for k in labels if str(k).startswith("Koszt wody – ")) == 6
    kwh = labels["Prąd (cała chlewnia)"][1]
    assert abs(labels["Koszt prądu"][1] - kwh * 0.85) < 0.05
    assert abs(labels["KOSZT RAZEM"][1] - (labels["Koszt prądu"][1] + labels["Koszt wody razem"][1])) < 0.05


def test_monthly_send(sa):
    r = sa.post(f"{API}/monthly-report/send?client_id={CID}")
    assert r.status_code == 200, r.text
    assert "recipients" in r.json() or "skipped" in r.json()


def test_final_rates(sa):
    sa.put(f"{API}/rates?client_id={CID}", json={"kwh": 0.85, "m3": 6.5})
    assert sa.get(f"{API}/rates?client_id={CID}").json() == {"kwh": 0.85, "m3": 6.5}
