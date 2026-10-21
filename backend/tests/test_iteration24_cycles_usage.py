# Iteration 24: per-chamber water meters, usage range (+xlsx), rearing cycles (start/end/list), RBAC, events
import io
import os
import time
import uuid
from datetime import datetime
from zoneinfo import ZoneInfo
import pytest
import requests
from pymongo import MongoClient

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = f"{BASE_URL}/api"
H = {"Origin": BASE_URL}
KOW = "kowalski-a"
MP = "chlodnia-01"
TAG = uuid.uuid4().hex[:6]
PW = "Haslo1234!x"
TODAY = datetime.now(ZoneInfo("Europe/Warsaw")).date().isoformat()
mongo = MongoClient("mongodb://localhost:27017")["test_database"]


def login(email, pw):
    ss = requests.Session()
    ss.headers.update(H)
    r = ss.post(f"{API}/auth/login", json={"email": email, "password": pw})
    assert r.status_code == 200, r.text
    return ss


@pytest.fixture(scope="module")
def sa():
    s = login("admin@chlewnia.pl", "Wentylacja!2026")
    cl = s.get(f"{API}/clients").json()["clients"]
    s.kow = next(c["id"] for c in cl if c["name"] == "Ferma Kowalski")
    s.uids = []
    s.emails = {}
    for role in ("operator", "viewer"):
        email = f"test_{TAG}_{role}@kowalski.pl"
        r = s.post(f"{API}/users", json={"email": email, "password": PW, "role": role, "client_id": s.kow})
        assert r.status_code in (200, 201), r.text
        s.uids.append(r.json().get("id") or r.json().get("user", {}).get("id"))
        s.emails[role] = email
    mongo.cycles.delete_many({"device_id": KOW, "chamber": 0})
    yield s
    mongo.cycles.delete_many({"device_id": KOW, "started_by": {"$regex": f"test_{TAG}"}})
    for uid in s.uids:
        s.delete(f"{API}/users/{uid}")


@pytest.fixture(scope="module")
def op(sa):
    return login(sa.emails["operator"], PW)


@pytest.fixture(scope="module")
def viewer(sa):
    return login(sa.emails["viewer"], PW)


# State: chambers include water fields
def test_state_has_water(sa):
    devs = {d["device_id"]: d for d in sa.get(f"{API}/devices").json()["devices"]}
    ch = devs[KOW]["state"]["chambers"][0]
    assert "water_total_l" in ch and "water_lpm" in ch, ch


# Usage range: structure + growth
def test_usage_range_grows(viewer):
    u = f"{API}/devices/{KOW}/usage/range?start=2026-01-01&end=2026-12-31"
    r1 = viewer.get(u)
    assert r1.status_code == 200, r1.text
    a = r1.json()
    assert len(a["water_l"]) == 6 and isinstance(a["kwh"], (int, float))
    assert abs(a["water_total_l"] - sum(a["water_l"])) < 0.5
    assert any(d["day"] == TODAY for d in a["days"])
    time.sleep(35)
    b = viewer.get(u).json()
    assert b["water_total_l"] > a["water_total_l"], (a["water_total_l"], b["water_total_l"])
    assert b["kwh"] >= a["kwh"]


@pytest.mark.parametrize("s,e", [("2026-13-01", "2026-12-31"), ("2026-05-01", "2026-04-01"), ("abc", "2026-01-01")])
def test_usage_range_invalid(viewer, s, e):
    assert viewer.get(f"{API}/devices/{KOW}/usage/range?start={s}&end={e}").status_code == 422


def test_usage_range_other_farm_404(viewer):
    assert viewer.get(f"{API}/devices/{MP}/usage/range?start=2026-01-01&end=2026-12-31").status_code == 404
    assert viewer.get(f"{API}/devices/{MP}/usage/range/xlsx?start=2026-01-01&end=2026-12-31").status_code == 404


# XLSX export
def test_usage_xlsx(viewer, sa):
    from openpyxl import load_workbook
    r = viewer.get(f"{API}/devices/{KOW}/usage/range/xlsx?start=2026-01-01&end={TODAY}")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Podsumowanie", "Dni"]
    names = [c["name"] if isinstance(c, dict) else c for c in
             {d["device_id"]: d for d in sa.get(f"{API}/devices").json()["devices"]}[KOW].get("chamber_names", [])]
    txt = " ".join(str(c.value) for row in wb["Podsumowanie"].iter_rows() for c in row if c.value is not None)
    hdr = " ".join(str(c.value) for c in wb["Dni"][1])
    assert "Woda razem" in txt and "Prąd" in txt
    for n in names:
        if n:
            assert n in txt and n in hdr
    assert wb["Dni"].max_row >= 2


# Cycles
def test_cycles_flow(op, viewer):
    base = f"{API}/devices/{KOW}"
    # viewer forbidden
    assert viewer.post(f"{base}/chambers/0/cycles", json={"start": "2026-10-01", "animals": 500}).status_code == 403
    # future
    assert op.post(f"{base}/chambers/0/cycles", json={"start": "2099-01-01"}).status_code == 422
    assert op.post(f"{base}/chambers/0/cycles", json={"start": "2026-02-30"}).status_code == 422
    assert op.post(f"{base}/chambers/9/cycles", json={"start": "2026-10-01"}).status_code == 404
    r = op.post(f"{base}/chambers/0/cycles", json={"start": "2026-10-01", "animals": 500})
    assert r.status_code == 200, r.text
    c = r.json()
    assert c["started_by"].startswith(f"test_{TAG}_operator") and c["animals"] == 500 and c["end"] is None
    assert c["days"] == (datetime.fromisoformat(TODAY) - datetime.fromisoformat("2026-10-01")).days + 1
    cid = c["id"]
    assert op.post(f"{base}/chambers/0/cycles", json={"start": "2026-10-05"}).status_code == 409
    lst = viewer.get(f"{base}/cycles?chamber=0").json()["cycles"]
    row = next(x for x in lst if x["id"] == cid)
    assert row["water_l"] > 0 and "kwh" in row
    # viewer cannot end
    assert viewer.post(f"{base}/cycles/{cid}/end", json={}).status_code == 403
    assert op.post(f"{base}/cycles/badid/end", json={}).status_code == 404
    e = op.post(f"{base}/cycles/{cid}/end", json={})
    assert e.status_code == 200, e.text
    ej = e.json()
    assert ej["end"] == TODAY and ej["ended_by"].startswith(f"test_{TAG}_operator")
    assert op.post(f"{base}/cycles/{cid}/end", json={}).status_code == 409
    ev = [x["message"] for x in viewer.get(f"{API}/events?device_id={KOW}&limit=50").json()["events"]]
    em = f"test_{TAG}_operator@kowalski.pl"
    assert any(f"nowe wstawienie 2026-10-01, 500 szt. ({em})" in m for m in ev), ev[:5]
    assert any("koniec chowu 2026-10-01" in m and em in m for m in ev), ev[:5]


def test_cycle_other_farm_404(op):
    assert op.post(f"{API}/devices/{MP}/chambers/0/cycles", json={"start": "2026-10-01"}).status_code == 404
    assert op.get(f"{API}/devices/{MP}/cycles?chamber=0").status_code == 404
