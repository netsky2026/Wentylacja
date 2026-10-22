"""Multi-tenant Fatek PLC MQTT gateway for piggery ventilation (clients -> PLC devices -> chambers)."""
from fastapi import FastAPI, APIRouter, HTTPException, WebSocket, WebSocketDisconnect, Depends, Request, UploadFile, File
from fastapi.responses import JSONResponse, StreamingResponse, Response
import io
import requests
from PIL import Image
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import re
import json
import time
import math
import copy
import socket
import secrets
import asyncio
import logging
import threading
import random
import uuid
import ipaddress
from zoneinfo import ZoneInfo
from pathlib import Path
from contextlib import asynccontextmanager
from datetime import datetime, timezone, timedelta
from typing import Any, Optional, List

import paho.mqtt.client as mqtt
from bson import ObjectId
from pymongo.errors import DuplicateKeyError
from cryptography.fernet import Fernet, InvalidToken
from pydantic import BaseModel, Field

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

import auth  # noqa: E402
import notify  # noqa: E402
from auth import require  # noqa: E402

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger("fatek")

MONGO_URL = os.environ['MONGO_URL']
DB_NAME = os.environ['DB_NAME']
FRONTEND_URL = os.environ["FRONTEND_URL"].rstrip("/")
MQTT_HOST = os.environ['MQTT_HOST']
MQTT_PORT = int(os.environ['MQTT_PORT'])
MQTT_PREFIX = os.environ['MQTT_PREFIX'].strip('/')
MQTT_USER = os.environ.get("MQTT_USER", "")
MQTT_PASS = os.environ.get("MQTT_PASS", "")
MQTT_TLS = False
SAMPLE_S = 15
ALARM_REMIND_S = 30 * 60
DEFAULT_TH = {"temp_max": 28.0, "co2_max": 3000.0}
LEGACY_DEVICE = "chlodnia-01"
FERNET = Fernet(os.environ["BROKER_SECRET"].encode())

LOGO_DIR = Path(os.environ["LOGO_DIR"])
STORAGE_APP = "marco-polo-wentylacja"


def put_object(path: str, data: bytes) -> dict:
    f = LOGO_DIR / path
    f.parent.mkdir(parents=True, exist_ok=True)
    f.write_bytes(data)
    return {"path": path}


def get_object(path: str) -> bytes:
    f = LOGO_DIR / path
    if not f.is_file():
        raise HTTPException(404, "No logo")
    return f.read_bytes()


THEME_RX = r"^(green|light|blue|amber|red|teal|graphite|lightblue|forest|grass|sun|orange)$"

client = AsyncIOMotorClient(MONGO_URL)
db = client[DB_NAME]

DEVICES: dict = {}     # device_id -> {client_id, client_name, name, chambers, simulated}
THRESHOLDS: dict = {}  # (device_id, chamber) -> {temp_max, co2_max}


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def state_topic(device_id: str) -> str:
    return f"{MQTT_PREFIX}/{device_id}/state"


def command_topic(device_id: str) -> str:
    return f"{MQTT_PREFIX}/{device_id}/command"


def make_client(tag: str) -> mqtt.Client:
    c = mqtt.Client(callback_api_version=mqtt.CallbackAPIVersion.VERSION2, client_id=f"{tag}-{uuid.uuid4().hex[:8]}")
    if MQTT_USER:
        c.username_pw_set(MQTT_USER, MQTT_PASS or None)
    if MQTT_TLS:
        c.tls_set()
    return c


def encrypt_secret(v: str) -> str:
    return FERNET.encrypt(v.encode()).decode() if v else ""


def decrypt_secret(v: str) -> str:
    try:
        return FERNET.decrypt(v.encode()).decode() if v else ""
    except InvalidToken:
        logger.warning("secret could not be decrypted")
        return ""


def th_for(device_id: str, chamber: int) -> dict:
    return THRESHOLDS.get((device_id, chamber), DEFAULT_TH)


# ============================================================
# Simulated Fatek PLC (one per simulated device, own MQTT client)
# ============================================================
class PiggerySimulator:
    def __init__(self, device_id: str, n: int):
        self.device_id = device_id
        self.enabled = True
        rnd = random.Random(device_id)
        self.chambers = []
        for i in range(n):
            sp = 20.0 + (i % 3) * 2
            self.chambers.append({
                "id": i, "temp": round(sp + rnd.uniform(-1.5, 2.0), 1), "setpoint": sp, "humidity": 65.0 + rnd.uniform(0, 8),
                "co2": 900.0 + rnd.uniform(0, 400), "on": True, "heating": False, "mode": "auto",
                "curtain_left": 30.0, "curtain_right": 30.0, "chimney": 40.0,
                "manual": {"curtain_left": 30.0, "curtain_right": 30.0, "chimney": 40.0},
                "water_lpm": 0.0, "water_total_l": round(rnd.uniform(500, 3000), 1)})
        self.plant = {"power_kw": 0.0, "power_kwh": round(rnd.uniform(50, 200), 1), "water_lpm": 0.0,
                      "water_total_l": round(rnd.uniform(2000, 8000), 1)}
        self._stop = threading.Event()
        self._thread: Optional[threading.Thread] = None
        self.mc: Optional[mqtt.Client] = None

    def mqtt_start(self):
        self.mc = make_client(f"sim-{self.device_id}"[:20])
        self.mc.on_connect = lambda c, u, f, rc, p: c.subscribe(command_topic(self.device_id), qos=1)
        self.mc.on_message = self._on_message
        self.mc.connect_async(MQTT_HOST, MQTT_PORT, keepalive=60)
        self.mc.loop_start()

    def mqtt_stop(self):
        try:
            self.mc.disconnect()
            self.mc.loop_stop()
        except Exception:
            pass

    def _on_message(self, c, u, msg):
        try:
            cmd = json.loads(msg.payload.decode())
        except Exception:
            return
        self.apply_command(cmd)
        self.publish()

    def apply_command(self, cmd: dict):
        t = cmd.get("type")
        try:
            ci = int(cmd.get("chamber", -1))
        except (TypeError, ValueError):
            return
        if not (0 <= ci < len(self.chambers)):
            return
        ch = self.chambers[ci]
        if t == "set_setpoint" and isinstance(cmd.get("value"), (int, float)):
            ch["setpoint"] = max(5.0, min(35.0, float(cmd["value"])))
        elif t == "set_mode" and cmd.get("value") in ("auto", "manual"):
            ch["mode"] = cmd["value"]
            if ch["mode"] == "manual":
                ch["manual"] = {k: ch[k] for k in VENT_FIELDS}
        elif t == "set_vent":
            for k, v in (cmd.get("value") or {}).items():
                if k in VENT_FIELDS and isinstance(v, (int, float)) and not isinstance(v, bool):
                    ch["manual"][k] = max(0.0, min(100.0, float(v)))
        elif t == "chamber_power":
            ch["on"] = bool(cmd.get("value"))
            if not ch["on"]:
                ch["heating"] = False

    @staticmethod
    def _approach(cur, target, rate):
        return min(target, cur + rate) if cur < target else max(target, cur - rate)

    def _step(self):
        for ch in self.chambers:
            if not ch["on"]:
                ch["temp"] += (18.0 - ch["temp"]) * 0.02
                for k in VENT_FIELDS:
                    ch[k] = self._approach(ch[k], 0.0, 4)
                ch["heating"] = False
                self._sensors(ch)
                continue
            err = ch["temp"] - ch["setpoint"]
            target = max(0.0, min(100.0, (err + 1.0) * 28.0))
            if ch["mode"] == "manual":
                tl, tr, tc = (ch["manual"][k] for k in VENT_FIELDS)
            else:
                tl, tr, tc = target, min(100.0, target + 4), target
            ch["curtain_left"] = round(self._approach(ch["curtain_left"], tl, 3), 1)
            ch["curtain_right"] = round(self._approach(ch["curtain_right"], tr, 3), 1)
            ch["chimney"] = round(self._approach(ch["chimney"], tc, 4), 1)
            ch["heating"] = err < -1.0
            vent = (ch["curtain_left"] + ch["curtain_right"] + ch["chimney"]) / 3.0
            ch["temp"] += 0.05 - vent * 0.004 + (0.18 if ch["heating"] else 0) + random.uniform(-0.05, 0.05)
            ch["temp"] = round(ch["temp"], 1)
            self._sensors(ch)
        fans = sum(ch["chimney"] for ch in self.chambers if ch["on"]) / 100.0
        heat = sum(1 for ch in self.chambers if ch["heating"])
        power = max(0.0, 1.2 + fans * 2.5 + heat * 3.0 + random.uniform(-0.2, 0.2))
        self.plant["power_kw"] = round(power, 2)
        self.plant["power_kwh"] = round(self.plant["power_kwh"] + power / 3600.0, 3)
        for ch in self.chambers:
            ch["water_lpm"] = round(random.uniform(0.6, 2.2), 2) if ch["on"] else 0.0
            ch["water_total_l"] = round(ch["water_total_l"] + ch["water_lpm"] / 60.0, 2)
        lpm = sum(ch["water_lpm"] for ch in self.chambers)
        self.plant["water_lpm"] = round(lpm, 2)
        self.plant["water_total_l"] = round(self.plant["water_total_l"] + lpm / 60.0, 2)

    @staticmethod
    def _sensors(ch):
        vent = (ch["curtain_left"] + ch["curtain_right"] + ch["chimney"]) / 3.0 if ch["on"] else 0.0
        hum_t = (78.0 - vent * 0.3) if ch["on"] else 60.0
        ch["humidity"] = round(max(35.0, min(99.0, ch["humidity"] + (hum_t - ch["humidity"]) * 0.03 + random.uniform(-0.3, 0.3))), 1)
        co2_t = (2400.0 - vent * 18.0) if ch["on"] else 500.0
        ch["co2"] = round(max(400.0, min(6000.0, ch["co2"] + (co2_t - ch["co2"]) * 0.03 + random.uniform(-10, 10))), 0)

    def snapshot(self) -> dict:
        return {"device_id": self.device_id, "model": "Fatek FBs-24MC (symulator)",
                "chambers": copy.deepcopy(self.chambers), "plant": dict(self.plant), "ts": now_iso()}

    def publish(self):
        if not self.enabled or not self.mc:
            return
        try:
            self.mc.publish(state_topic(self.device_id), json.dumps(self.snapshot()), qos=1)
        except Exception as e:
            logger.warning("SIM publish failed: %s", e)

    def _loop(self):
        while not self._stop.is_set():
            self._step()
            self.publish()
            self._stop.wait(1.0)

    def start(self):
        self.mqtt_start()
        self._thread = threading.Thread(target=self._loop, daemon=True)
        self._thread.start()

    def stop(self):
        self._stop.set()
        self.mqtt_stop()


class SimManager:
    def __init__(self):
        self.sims: dict = {}

    def sync(self):
        wanted = {d: v["chambers"] for d, v in DEVICES.items() if v["simulated"]}
        for d in list(self.sims):
            if d not in wanted or len(self.sims[d].chambers) != wanted[d]:
                self.sims.pop(d).stop()
        for d, n in wanted.items():
            if d not in self.sims:
                self.sims[d] = PiggerySimulator(d, n)
                self.sims[d].start()

    def restart_mqtt(self):
        for s in self.sims.values():
            s.mqtt_stop()
            s.mqtt_start()

    def stop_all(self):
        for s in self.sims.values():
            s.stop()
        self.sims.clear()


# ============================================================
# Payload validation
# ============================================================
NUM_FIELDS = ("temp", "setpoint", "humidity", "co2", "curtain_left", "curtain_right", "chimney")
VENT_FIELDS = ("curtain_left", "curtain_right", "chimney")
PLANT_FIELDS = ("power_kw", "power_kwh", "water_lpm", "water_total_l")
CH_OPT_FIELDS = ("water_lpm", "water_total_l")


def _num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


def ch_name(device_id: str, cid: int) -> str:
    d = DEVICES.get(device_id) or {}
    return d.get("chamber_names", {}).get(str(cid)) or f"Komora {cid + 1}"


def sanitize_state(p, device_id: str) -> Optional[dict]:
    """Accept only well-formed numeric state with exactly the device's chambers; drop anything else."""
    n = DEVICES[device_id]["chambers"]
    if not isinstance(p, dict) or not isinstance(p.get("chambers"), list):
        return None
    ids = [c.get("id") if isinstance(c, dict) else None for c in p["chambers"]]
    if len(ids) != n or not all(_num(i) for i in ids) or sorted(int(i) for i in ids) != list(range(n)):
        return None
    chambers = []
    for ch in sorted(p["chambers"], key=lambda c: int(c["id"])):
        cid = int(ch["id"])
        out = {"id": cid, "name": ch_name(device_id, cid), "on": bool(ch.get("on")), "heating": bool(ch.get("heating")),
               "mode": "manual" if ch.get("mode") == "manual" else "auto"}
        for k in NUM_FIELDS:
            if not _num(ch.get(k)):
                return None
            out[k] = float(ch[k])
        for k in CH_OPT_FIELDS:
            if _num(ch.get(k)):
                out[k] = float(ch[k])
        man = ch.get("manual")
        if isinstance(man, dict):
            out["manual"] = {k: float(man[k]) for k in VENT_FIELDS if _num(man.get(k))}
        chambers.append(out)
    plant_in = p.get("plant") if isinstance(p.get("plant"), dict) else {}
    return {"device_id": device_id, "model": str(p.get("model", ""))[:64], "chambers": chambers,
            "plant": {k: float(plant_in[k]) for k in PLANT_FIELDS if _num(plant_in.get(k))}, "ts": now_iso()}


def state_error(p, device_id: str) -> str:
    """Human-readable reason why sanitize_state rejected a message (MQTT preview)."""
    n = DEVICES[device_id]["chambers"]
    if not isinstance(p, dict):
        return "Wiadomość musi być obiektem JSON { ... }"
    if not isinstance(p.get("chambers"), list):
        return "Brak listy \"chambers\""
    if len(p["chambers"]) != n:
        return f"Zła liczba komór: przysłano {len(p['chambers'])}, sterownik ma {n}"
    ids = [c.get("id") if isinstance(c, dict) else None for c in p["chambers"]]
    if not all(_num(i) for i in ids) or sorted(int(i) for i in ids) != list(range(n)):
        return f"Pola \"id\" komór muszą być liczbami 0–{n - 1} (przysłano {ids})"
    for ch in p["chambers"]:
        bad = [k for k in NUM_FIELDS if not _num(ch.get(k))]
        if bad:
            return f"Komora id={ch['id']}: brak lub nieprawidłowa wartość pól: {', '.join(bad)} (muszą być liczbami)"
    return "Nieprawidłowy format wiadomości"


def annotate(device_id: str, payload: dict) -> dict:
    for ch in payload["chambers"]:
        th = th_for(device_id, ch["id"])
        ch["temp_max"], ch["co2_max"] = th["temp_max"], th["co2_max"]
        ch["temp_alarm"] = bool(ch["on"] and ch["temp"] > th["temp_max"])
        ch["co2_alarm"] = bool(ch["co2"] > th["co2_max"])
        ch["alarm"] = ch["temp_alarm"] or ch["co2_alarm"]
    return payload


def alarm_reason(ch: dict) -> str:
    r = []
    if ch.get("temp_alarm"):
        r.append(f"temperatura {ch['temp']}°C > {ch.get('temp_max')}°C")
    if ch.get("co2_alarm"):
        r.append(f"CO₂ {int(ch['co2'])} ppm > {int(ch.get('co2_max'))} ppm")
    return ", ".join(r)


def dev_label(device_id: str) -> str:
    d = DEVICES.get(device_id, {})
    return f"{d.get('client_name', '')} / {d.get('name', device_id)}"


# ============================================================
# Gateway
# ============================================================
class Gateway:
    def __init__(self):
        self.loop: Optional[asyncio.AbstractEventLoop] = None
        self.ws_clients: dict = {}  # ws -> user
        self.states: dict = {}
        self.last_seen: dict = {}
        self.raw: dict = {}
        self.last_cmd: dict = {}
        self.connected = False
        self.mc: Optional[mqtt.Client] = None

    def mqtt_start(self):
        self.mc = make_client("fatek-api")
        self.mc.on_connect = self._on_connect
        self.mc.on_disconnect = self._on_disconnect
        self.mc.on_message = self._on_message
        self.mc.connect_async(MQTT_HOST, MQTT_PORT, keepalive=60)
        self.mc.loop_start()

    def mqtt_stop(self):
        try:
            self.mc.disconnect()
            self.mc.loop_stop()
        except Exception:
            pass
        self.connected = False

    def _on_connect(self, c, u, flags, rc, props):
        self.connected = True
        c.subscribe(f"{MQTT_PREFIX}/+/state", qos=1)
        self.schedule(log_event("info", "Połączono z brokerem MQTT", {"host": MQTT_HOST}))

    def _on_disconnect(self, c, u, flags, rc, props):
        self.connected = False

    def _on_message(self, c, u, msg):
        parts = msg.topic.split("/")
        if len(parts) < 2 or "/".join(parts[:-2]) != MQTT_PREFIX or parts[-1] != "state":
            return
        device_id = parts[-2]
        if device_id not in DEVICES:
            return
        raw = msg.payload[:8192].decode("utf-8", "replace")
        rec = {"topic": msg.topic, "ts": now_iso(), "size": len(msg.payload), "payload": raw, "ok": False, "error": None}
        self.raw[device_id] = rec
        try:
            data = json.loads(msg.payload.decode())
        except Exception:
            rec["error"] = "To nie jest poprawny JSON (UTF-8)"
            return
        payload = sanitize_state(data, device_id)
        if payload is None:
            rec["error"] = state_error(data, device_id)
            return
        rec["ok"] = True
        prev = self.states.get(device_id)
        annotate(device_id, payload)
        self.states[device_id] = payload
        self.last_seen[device_id] = time.time()
        self.detect_alarms(device_id, prev, payload)
        self.schedule(broadcast({"type": "state", "payload": payload}, DEVICES[device_id]["client_id"]))

    def detect_alarms(self, device_id, prev, cur):
        was = {c["id"]: c.get("alarm") for c in (prev or {}).get("chambers", [])}
        for ch in cur["chambers"]:
            if ch["alarm"] and not was.get(ch["id"]):
                title = f"ALARM: {dev_label(device_id)} – {ch_name(device_id, ch['id'])}"
                reason = alarm_reason(ch)
                self.schedule(log_event("warning", f"{title} — {reason}", {"chamber": ch["id"]}, device_id))
                self.schedule(broadcast({"type": "alarm", "device_id": device_id, "title": title, "body": reason},
                                        DEVICES[device_id]["client_id"]))

    def schedule(self, coro):
        if self.loop and not self.loop.is_closed():
            asyncio.run_coroutine_threadsafe(coro, self.loop)
        else:
            coro.close()

    def publish_command(self, device_id: str, cmd: dict) -> bool:
        ok = self.mc.publish(command_topic(device_id), json.dumps(cmd), qos=1).rc == mqtt.MQTT_ERR_SUCCESS
        self.last_cmd[device_id] = {"topic": command_topic(device_id), "ts": now_iso(), "payload": json.dumps(cmd, ensure_ascii=False), "ok": ok}
        return ok

    def is_online(self, device_id: str) -> bool:
        t = self.last_seen.get(device_id)
        return t is not None and time.time() - t < 10


gateway = Gateway()
sims = SimManager()


async def broadcast(data: dict, client_id: Optional[str]):
    for ws, user in list(gateway.ws_clients.items()):
        if user["role"] != "superadmin" and user.get("client_id") != client_id:
            continue
        try:
            await ws.send_json(data)
        except Exception:
            gateway.ws_clients.pop(ws, None)


async def log_event(level: str, message: str, meta: dict, device_id: Optional[str] = None,
                    client_id: Optional[str] = None):
    if device_id and not client_id:
        client_id = DEVICES.get(device_id, {}).get("client_id")
    doc = {"level": level, "message": message, "meta": meta, "device_id": device_id, "client_id": client_id,
           "ts": now_iso()}
    await db.plc_events.insert_one(doc)
    doc.pop("_id", None)
    if client_id:
        await broadcast({"type": "event", "event": doc}, client_id)
    else:
        for ws, user in list(gateway.ws_clients.items()):
            if user["role"] == "superadmin":
                try:
                    await ws.send_json({"type": "event", "event": doc})
                except Exception:
                    gateway.ws_clients.pop(ws, None)


# ============================================================
# Registry, migration, background tasks
# ============================================================
async def load_registry():
    cdocs = {str(c["_id"]): c async for c in db.clients.find({})}
    clients = {k: c["name"] for k, c in cdocs.items()}
    DEVICES.clear()
    async for d in db.devices.find({}):
        DEVICES[d["device_id"]] = {"client_id": d["client_id"], "client_name": clients.get(d["client_id"], "?"),
                                   "name": d["name"], "chambers": d["chambers"], "simulated": d.get("simulated", False),
                                   "chamber_names": d.get("chamber_names", {}),
                                   "auto_return": cdocs.get(d["client_id"], {}).get("auto_return", "00:00")}
    THRESHOLDS.clear()
    async for t in db.device_thresholds.find({}):
        THRESHOLDS[(t["device_id"], t["chamber"])] = {"temp_max": t["temp_max"], "co2_max": t["co2_max"]}
    for d in list(gateway.states):
        if d not in DEVICES or len(gateway.states[d]["chambers"]) != DEVICES[d]["chambers"]:
            gateway.states.pop(d, None)


async def migrate_legacy():
    if await db.devices.count_documents({}) or await db.clients.count_documents({}):
        return
    res = await db.clients.insert_one({"name": "Marco Polo", "created_at": now_iso()})
    cid = str(res.inserted_id)
    await db.devices.insert_one({"client_id": cid, "name": "Sterownik 1", "device_id": LEGACY_DEVICE, "chambers": 3,
                                 "simulated": True, "mqtt_username": LEGACY_DEVICE, "mqtt_password_enc": "",
                                 "created_at": now_iso()})
    async for t in db.chamber_thresholds.find({}):
        await db.device_thresholds.update_one({"device_id": LEGACY_DEVICE, "chamber": t["chamber"]},
                                              {"$set": {"temp_max": t["temp_max"], "co2_max": t["co2_max"]}}, upsert=True)
    for coll in ("plc_history", "plc_plant_history", "plc_events"):
        await db[coll].update_many({"device_id": {"$exists": False}}, {"$set": {"device_id": LEGACY_DEVICE}})
    await db.plc_events.update_many({"client_id": {"$exists": False}}, {"$set": {"client_id": cid}})
    alarm = await db.alarm_settings.find_one({"_id": "main"})
    if alarm:
        await db.alarm_settings.insert_one({**alarm, "_id": cid})
    await db.users.update_many({"role": {"$ne": "superadmin"}, "client_id": {"$exists": False}}, {"$set": {"client_id": cid}})


def set_broker_globals(cfg: dict):
    global MQTT_HOST, MQTT_PORT, MQTT_USER, MQTT_PASS, MQTT_TLS, MQTT_PREFIX
    MQTT_HOST, MQTT_PORT = cfg["host"], int(cfg["port"])
    MQTT_USER, MQTT_PASS = cfg.get("username") or "", cfg.get("password") or ""
    MQTT_TLS, MQTT_PREFIX = bool(cfg.get("tls")), cfg["prefix"].strip("/")


async def load_broker():
    doc = await db.broker_config.find_one({"_id": "main"})
    if doc:
        doc["password"] = decrypt_secret(doc.get("password_enc", ""))
        set_broker_globals(doc)


MQTT_DYNSEC = os.environ.get("MQTT_DYNSEC") == "1"  # self-hosted Mosquitto with dynamic-security plugin


def dynsec(commands: list):
    """Provision per-device broker accounts (Mosquitto dynamic security). No-op on external brokers."""
    if MQTT_DYNSEC and gateway.mc:
        gateway.mc.publish("$CONTROL/dynamic-security/v1", json.dumps({"commands": commands}), qos=1)


def dynsec_add_device(device_id: str, password: str):
    role = f"dev-{device_id}"
    dynsec([
        {"command": "createRole", "rolename": role, "acls": [
            {"acltype": "publishClientSend", "topic": state_topic(device_id), "allow": True},
            {"acltype": "subscribePattern", "topic": command_topic(device_id), "allow": True},
            {"acltype": "publishClientReceive", "topic": command_topic(device_id), "allow": True}]},
        {"command": "deleteClient", "username": device_id},
        {"command": "createClient", "username": device_id, "password": password, "roles": [{"rolename": role}]},
    ])


def dynsec_remove_device(device_id: str):
    dynsec([{"command": "deleteClient", "username": device_id}, {"command": "deleteRole", "rolename": f"dev-{device_id}"}])


def is_public_host(host: str, port: int) -> bool:
    try:
        infos = socket.getaddrinfo(host, port, proto=socket.IPPROTO_TCP)
    except socket.gaierror:
        return False
    return bool(infos) and all(ipaddress.ip_address(i[4][0]).is_global for i in infos)


async def history_sampler():
    while True:
        await asyncio.sleep(SAMPLE_S)
        ts = now_iso()
        for device_id, st in list(gateway.states.items()):
            if not gateway.is_online(device_id):
                continue
            docs = [{"device_id": device_id, "chamber": ch["id"], **{k: ch[k] for k in NUM_FIELDS if k != "setpoint"},
                     "ts": ts} for ch in st["chambers"]]
            try:
                await db.plc_history.insert_many(docs)
                if st.get("plant"):
                    await db.plc_plant_history.insert_one({"device_id": device_id, "power_kw": st["plant"].get("power_kw"),
                                                           "water_lpm": st["plant"].get("water_lpm"), "ts": ts})
            except Exception as e:
                logger.warning("history insert failed: %s", e)
            try:
                await accumulate_meters(device_id, st)
            except Exception as e:
                logger.warning("meter accumulate failed: %s", e)


METER_MAX_STEP = {"power": 1000.0, "water": 1_000_000.0}


def local_day(dt: Optional[datetime] = None) -> str:
    return (dt or datetime.now(timezone.utc)).astimezone(ZoneInfo("Europe/Warsaw")).date().isoformat()


async def accumulate_meters(device_id: str, st: dict):
    """Adds the increase of cumulative meters (plant kWh, per-chamber water litres) to the Warsaw-local daily doc."""
    cur = {"power": st.get("plant", {}).get("power_kwh")}
    for ch in st["chambers"]:
        cur[f"w{ch['id']}"] = ch.get("water_total_l")
    cur = {k: v for k, v in cur.items() if v is not None}
    if not cur:
        return
    last = await db.meter_last.find_one({"_id": device_id}) or {}
    inc = {}
    for k, v in cur.items():
        d = v - last[k] if k in last else None
        if d is not None and 0 < d <= METER_MAX_STEP["power" if k == "power" else "water"]:
            inc["kwh" if k == "power" else f"water.{k[1:]}"] = round(d, 4)
    await db.meter_last.update_one({"_id": device_id}, {"$set": cur}, upsert=True)
    if inc:
        day = local_day()
        await db.meter_daily.update_one({"_id": f"{device_id}|{day}"}, {"$inc": inc, "$set": {"device_id": device_id, "day": day}},
                                        upsert=True)


async def usage_between(device_id: str, start: str, end: str) -> dict:
    rows = await db.meter_daily.find({"device_id": device_id, "day": {"$gte": start, "$lte": end}}).sort("day", 1).to_list(2000)
    n = DEVICES[device_id]["chambers"]
    water = [round(sum(r.get("water", {}).get(str(i), 0) for r in rows), 1) for i in range(n)]
    return {"from": start, "to": end, "kwh": round(sum(r.get("kwh", 0) for r in rows), 2), "water_l": water,
            "water_total_l": round(sum(water), 1),
            "days": [{"day": r["day"], "kwh": round(r.get("kwh", 0), 2),
                      "water_l": [round(r.get("water", {}).get(str(i), 0), 1) for i in range(n)]} for r in rows]}


alarm_active: dict = {}  # "device:chamber" -> last notification epoch


async def ws_revalidator():
    """Drop WS sessions whose user was deleted or changed role/client/password."""
    while True:
        await asyncio.sleep(30)
        for ws, u in list(gateway.ws_clients.items()):
            cur = await db.users.find_one({"_id": u["_id"]})
            if (not cur or any(cur.get(k) != u.get(k) for k in ("role", "client_id", "tv"))
                    or await auth.session_revoked(u.get("_sid"))):
                gateway.ws_clients.pop(ws, None)
                try:
                    await ws.close(code=4401)
                except Exception:
                    pass


async def save_alarm_state():
    await db.alarm_state.update_one({"_id": "main"}, {"$set": {"active": alarm_active}}, upsert=True)


async def send_alarm(client_id: str, subject: str, body: str, kind: str):
    cfg = await db.alarm_settings.find_one({"_id": client_id}) or {}
    recipients = cfg.get("recipients", []) if cfg.get("enabled", True) else []
    if not recipients:
        return
    ts = now_iso()
    results = await asyncio.gather(*[asyncio.to_thread(notify.deliver, r["type"], r["value"], subject, body)
                                     for r in recipients])
    await db.alarm_notifications.insert_many([
        {"client_id": client_id, "ts": ts, "kind": kind, "channel": r["type"], "to": r["value"], "subject": subject,
         "body": body, "status": st, "error": err} for r, (st, err) in zip(recipients, results)])
    sent = sum(st == "sent" for st, _ in results)
    failed = [r["value"] for r, (st, _) in zip(recipients, results) if st == "failed"]
    logger.info("ALARM NOTIFY %s -> %s", subject, [(r["value"], st) for r, (st, _) in zip(recipients, results)])
    await log_event("warning" if failed else "info",
                    f"Powiadomienie: wysłano {sent}/{len(recipients)}" + (f", błąd: {', '.join(failed)}" if failed else "")
                    + f" – {subject}", {"kind": kind}, client_id=client_id)


async def alarm_notifier():
    doc = await db.alarm_state.find_one({"_id": "main"}) or {}
    alarm_active.update({k: v for k, v in doc.get("active", {}).items() if ":" in k})
    while True:
        await asyncio.sleep(5)
        now = time.time()
        try:
            for device_id, st in list(gateway.states.items()):
                if not gateway.is_online(device_id) or device_id not in DEVICES:
                    continue
                cid = DEVICES[device_id]["client_id"]
                for ch in st["chambers"]:
                    key, name = f"{device_id}:{ch['id']}", f"{dev_label(device_id)} – {ch_name(device_id, ch['id'])}"
                    last = alarm_active.get(key)
                    if ch.get("alarm"):
                        if last is None or now - last >= ALARM_REMIND_S:
                            alarm_active[key] = now
                            await save_alarm_state()
                            subj = f"ALARM: {name}" if last is None else f"PRZYPOMNIENIE – ALARM trwa: {name}"
                            await send_alarm(cid, subj, alarm_reason(ch), "alarm" if last is None else "reminder")
                    elif last is not None:
                        alarm_active.pop(key)
                        await save_alarm_state()
                        await send_alarm(cid, f"Koniec alarmu: {name}",
                                         f"Temp {ch['temp']}°C, CO₂ {int(ch['co2'])} ppm", "cleared")
        except Exception as e:
            logger.warning("alarm notifier failed: %s", e)


PL_TZ = ZoneInfo("Europe/Warsaw")


async def reset_manual_to_auto(hm: Optional[str] = None):
    for device_id, st in list(gateway.states.items()):
        if hm and (DEVICES.get(device_id) or {}).get("auto_return") != hm:
            continue
        for ch in st.get("chambers", []):
            if ch.get("mode") == "manual" and gateway.publish_command(
                    device_id, {"type": "set_mode", "chamber": ch["id"], "value": "auto", "request_id": uuid.uuid4().hex}):
                await log_event("info", f"{dev_label(device_id)} – {ch_name(device_id, ch['id'])}: automatyczny powrót do AUTO ({hm or 'ręcznie'})",
                                {"user": "system"}, device_id)


async def midnight_auto():
    while True:
        now = datetime.now(timezone.utc)
        await asyncio.sleep(62 - now.second - now.microsecond / 1e6)
        try:
            await reset_manual_to_auto(datetime.now(PL_TZ).strftime("%H:%M"))
        except Exception as e:
            logger.error("midnight auto reset failed: %s", e)
        try:
            await monthly_tick()
        except Exception as e:
            logger.error("monthly tick failed: %s", e)


@asynccontextmanager
async def lifespan(app: FastAPI):
    gateway.loop = asyncio.get_running_loop()
    auth.init(db)
    await auth.setup()
    await migrate_legacy()
    await migrate_slugs()
    await db.clients.create_index("slug", unique=True, partialFilterExpression={"slug": {"$type": "string"}})
    await db.cycles.create_index([("device_id", 1), ("chamber", 1)], unique=True, name="one_active_cycle",
                                 partialFilterExpression={"active": True})
    await load_broker()
    await load_registry()
    await db.plc_history.create_index([("device_id", 1), ("ts", 1)])
    await db.plc_plant_history.create_index([("device_id", 1), ("ts", 1)])
    await db.devices.create_index("device_id", unique=True)
    await db.contact_requests.create_index("ts", expireAfterSeconds=3600)
    gateway.mqtt_start()
    sims.sync()
    tasks = [asyncio.create_task(history_sampler()), asyncio.create_task(alarm_notifier()),
             asyncio.create_task(ws_revalidator()), asyncio.create_task(midnight_auto())]
    yield
    for t in tasks:
        t.cancel()
    sims.stop_all()
    gateway.mqtt_stop()
    client.close()


app = FastAPI(lifespan=lifespan)
api = APIRouter(prefix="/api")


# ============================================================
# Access helpers
# ============================================================
def visible(user: dict, device_id: str) -> bool:
    d = DEVICES.get(device_id)
    return bool(d) and (user["role"] == "superadmin" or d["client_id"] == user.get("client_id"))


def device_or_404(user: dict, device_id: str) -> dict:
    if not visible(user, device_id):
        raise HTTPException(404, "Device not found")
    return DEVICES[device_id]


def client_scope(user: dict, client_id: Optional[str]) -> str:
    if user["role"] == "superadmin":
        if not client_id or not ObjectId.is_valid(client_id):
            raise HTTPException(422, "client_id required")
        return client_id
    return user.get("client_id")


def device_view(device_id: str, user: dict) -> dict:
    d = DEVICES[device_id]
    out = {"device_id": device_id, "name": d["name"], "client_id": d["client_id"], "client_name": d["client_name"],
           "chambers": d["chambers"], "chamber_names": d["chamber_names"], "simulated": d["simulated"],
           "auto_return": d["auto_return"], "online": gateway.is_online(device_id),
           "state": gateway.states.get(device_id)}
    if user["role"] == "superadmin":
        out.update({"state_topic": state_topic(device_id), "command_topic": command_topic(device_id)})
    return out


# ============================================================
# Models
# ============================================================
class CommandIn(BaseModel):
    type: str = Field(pattern=r"^(set_setpoint|set_mode|set_vent)$")
    chamber: int = Field(ge=0, le=9)
    value: Any = None


class ThresholdIn(BaseModel):
    temp_max: float = Field(ge=-20, le=60)
    co2_max: float = Field(ge=300, le=20000)


class BrokerIn(BaseModel):
    host: str = Field(min_length=1, max_length=255, pattern=r"^[A-Za-z0-9.\-]+$")
    port: int = Field(ge=1, le=65535)
    username: str = Field(default="", max_length=128)
    password: Optional[str] = Field(default=None, max_length=256)
    tls: bool = False
    prefix: str = Field(min_length=1, max_length=128, pattern=r"^[A-Za-z0-9_\-/]+$")


class ClientIn(BaseModel):
    name: str = Field(min_length=1, max_length=80)


class ClientPatch(BaseModel):
    name: Optional[str] = Field(default=None, min_length=1, max_length=80)
    display_name: Optional[str] = Field(default=None, max_length=40)
    slug: Optional[str] = Field(default=None, max_length=40, pattern=r"^([a-z0-9][a-z0-9-]{1,39})?$")
    theme: Optional[str] = Field(default=None, pattern=THEME_RX)
    contact: Optional["SiteIn"] = None


class DeviceIn(BaseModel):
    client_id: str
    name: str = Field(min_length=1, max_length=80)
    device_id: str = Field(min_length=3, max_length=48, pattern=r"^[a-z0-9][a-z0-9_\-]*$")
    chambers: int = Field(ge=1, le=10)
    simulated: bool = False


class DevicePatch(BaseModel):
    name: Optional[str] = Field(default=None, min_length=1, max_length=80)
    simulated: Optional[bool] = None


class Recipient(BaseModel):
    type: str = Field(pattern=r"^(email|sms)$")
    value: str = Field(min_length=3, max_length=254)


class AlarmSettingsIn(BaseModel):
    enabled: bool = True
    recipients: List[Recipient] = Field(default_factory=list, max_length=20)


EMAIL_RX = r"^[^@\s;,]+@[^@\s;,]+\.[^@\s;,]+$"
PHONE_RX = r"^\+?[0-9]{9,15}$"


# ============================================================
# Routes
# ============================================================
@api.get("/")
async def root():
    return {"message": "Fatek MQTT piggery gateway"}


@api.get("/health")
async def health(user: dict = Depends(require("viewer"))):
    return {"ok": True, "broker_connected": gateway.connected}


@api.get("/config")
async def get_config(user: dict = Depends(require("viewer"))):
    cfg = {"broker_host": MQTT_HOST, "broker_port": MQTT_PORT, "prefix": MQTT_PREFIX, "tls": MQTT_TLS}
    if user["role"] == "superadmin":
        cfg.update({"username": MQTT_USER, "has_password": bool(MQTT_PASS)})
    return cfg


@api.put("/broker")
async def set_broker(body: BrokerIn, user: dict = Depends(require("superadmin"))):
    cfg = body.model_dump()
    if not MQTT_DYNSEC and not await asyncio.to_thread(is_public_host, cfg["host"], cfg["port"]):
        raise HTTPException(422, "Broker host must resolve to a public address")
    if cfg["password"] is None:
        cfg["password"] = MQTT_PASS
    stored = {k: v for k, v in cfg.items() if k != "password"}
    stored["password_enc"] = encrypt_secret(cfg["password"])
    await db.broker_config.update_one({"_id": "main"}, {"$set": stored, "$unset": {"password": ""}}, upsert=True)
    await asyncio.to_thread(gateway.mqtt_stop)
    set_broker_globals(cfg)
    gateway.states.clear()
    gateway.last_seen.clear()
    gateway.mqtt_start()
    await asyncio.to_thread(sims.restart_mqtt)
    await log_event("info", f"Zmieniono broker: {MQTT_HOST}:{MQTT_PORT}", {"user": user["email"]})
    return await get_config(user)


@api.get("/clients")
async def list_clients(user: dict = Depends(require("superadmin"))):
    await migrate_slugs()
    out = []
    async for c in db.clients.find({}).sort("name", 1):
        cid = str(c["_id"])
        out.append({"id": cid, "name": c["name"], "devices": sum(1 for d in DEVICES.values() if d["client_id"] == cid),
                    "display_name": c.get("display_name", ""), "theme": c.get("theme", "green"), "slug": c.get("slug", ""),
                    "contact": client_contact(c), "logo_version": c.get("logo_version") if c.get("logo_path") else None})
    return {"clients": out}


PL_MAP = str.maketrans("ąćęłńóśźżĄĆĘŁŃÓŚŹŻ", "acelnoszzACELNOSZZ")


async def unique_slug(name: str) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", name.translate(PL_MAP).lower()).strip("-")[:34].strip("-")
    base = base if len(base) >= 2 else "ferma"
    slug, n = base, 2
    while await db.clients.find_one({"slug": slug}):
        slug, n = f"{base}-{n}", n + 1
    return slug


async def migrate_slugs():
    async for c in db.clients.find({"$or": [{"slug": {"$exists": False}}, {"slug": ""}, {"slug": None}]}):
        await db.clients.update_one({"_id": c["_id"]}, {"$set": {"slug": await unique_slug(c.get("name", ""))}})


@api.post("/clients")
async def create_client(body: ClientIn, user: dict = Depends(require("superadmin"))):
    name = body.name.strip()
    slug = await unique_slug(name)
    res = await db.clients.insert_one({"name": name, "slug": slug, "created_at": now_iso()})
    await log_event("info", f"Dodano klienta: {name}", {"user": user["email"]}, client_id=str(res.inserted_id))
    return {"id": str(res.inserted_id), "name": name, "slug": slug, "devices": 0}


async def client_or_404(client_id: str) -> dict:
    c = await db.clients.find_one({"_id": ObjectId(client_id)}) if ObjectId.is_valid(client_id) else None
    if not c:
        raise HTTPException(404, "Client not found")
    return c


@api.patch("/clients/{client_id}")
async def patch_client(client_id: str, body: ClientPatch, user: dict = Depends(require("superadmin"))):
    c = await client_or_404(client_id)
    upd = {k: (v.strip() if isinstance(v, str) else v) for k, v in body.model_dump().items() if v is not None}
    if "contact" in upd:
        upd["contact"] = {k: v.strip() for k, v in upd["contact"].items()}
    if upd.get("slug") == "" and c.get("slug"):
        upd.pop("slug")
    if upd.get("slug") and await db.clients.find_one({"slug": upd["slug"], "_id": {"$ne": c["_id"]}}):
        raise HTTPException(409, "Link name already used")
    if upd:
        await db.clients.update_one({"_id": c["_id"]}, {"$set": upd})
        await load_registry()
    return {"ok": True}


async def store_logo(file: UploadFile, folder: str) -> tuple:
    data = await file.read(1_000_001)
    if len(data) > 1_000_000:
        raise HTTPException(413, "Logo too large (max 1 MB)")
    try:
        img = Image.open(io.BytesIO(data))
        if img.format not in ("PNG", "JPEG", "WEBP"):
            raise ValueError("format")
        img = img.convert("RGBA")
        img.thumbnail((256, 256))
        buf = io.BytesIO()
        img.save(buf, "PNG")
    except Exception:
        raise HTTPException(422, "Unsupported image (PNG, JPEG or WEBP)")
    res = await asyncio.to_thread(put_object, f"{STORAGE_APP}/logos/{folder}/{uuid.uuid4().hex}.png", buf.getvalue())
    return res["path"], uuid.uuid4().hex[:8]


@api.post("/clients/{client_id}/logo")
async def upload_logo(client_id: str, file: UploadFile = File(...), user: dict = Depends(require("superadmin"))):
    c = await client_or_404(client_id)
    path, version = await store_logo(file, client_id)
    await db.clients.update_one({"_id": c["_id"]}, {"$set": {"logo_path": path, "logo_version": version}})
    return {"ok": True, "logo_version": version}


@api.delete("/clients/{client_id}/logo")
async def delete_logo(client_id: str, user: dict = Depends(require("superadmin"))):
    c = await client_or_404(client_id)
    await db.clients.update_one({"_id": c["_id"]}, {"$unset": {"logo_path": "", "logo_version": ""}})
    return {"ok": True}


@api.get("/clients/{client_id}/logo")
async def get_logo(client_id: str, user: dict = Depends(require("viewer"))):
    if user["role"] != "superadmin" and user.get("client_id") != client_id:
        raise HTTPException(404, "Not found")
    c = await client_or_404(client_id)
    if not c.get("logo_path"):
        raise HTTPException(404, "No logo")
    data = await asyncio.to_thread(get_object, c["logo_path"])
    return Response(content=data, media_type="image/png", headers={"Cache-Control": "private, max-age=3600"})


def _square_icon(png: bytes, size: int, bg: str) -> bytes:
    src = Image.open(io.BytesIO(png)).convert("RGBA")
    canvas = Image.new("RGBA", (size, size), bg)
    inner = int(size * 0.8)
    src = src.resize((int(src.width * inner / max(src.size)), int(src.height * inner / max(src.size))), Image.LANCZOS)
    canvas.paste(src, ((size - src.width) // 2, (size - src.height) // 2), src)
    out = io.BytesIO()
    canvas.convert("RGB").save(out, "PNG")
    return out.getvalue()


THEME_BG = {"green": "#0A0A0A", "light": "#F3F5F3", "blue": "#080B12", "amber": "#0C0A07", "red": "#0E0909", "teal": "#061010",
            "graphite": "#111214", "lightblue": "#F2F5F9", "forest": "#050F07", "grass": "#2E9A3E", "sun": "#FFD000", "orange": "#FF8A1F"}


async def client_by_slug(slug: str) -> dict:
    c = await db.clients.find_one({"slug": slug}) if re.match(r"^[a-z0-9][a-z0-9-]{1,39}$", slug or "") else None
    if not c:
        raise HTTPException(404, "Unknown client link")
    return c


def client_contact(c: dict) -> dict:
    ct = c.get("contact") or {}
    return {k: ct.get(k, "") for k in ("phone", "email", "address", "website")}


@api.get("/public/branding/{slug}")
async def public_branding(slug: str):
    c = await client_by_slug(slug)
    return {"slug": slug, "app_name": c.get("display_name") or None, "theme": c.get("theme", "green"), "contact": client_contact(c),
            "logo_version": c.get("logo_version") if c.get("logo_path") else None}


@api.get("/public/logo/{slug}")
async def public_logo(slug: str, size: int = 0):
    c = await client_by_slug(slug)
    if not c.get("logo_path"):
        raise HTTPException(404, "No logo")
    data = await asyncio.to_thread(get_object, c["logo_path"])
    if size in (180, 192, 512):
        data = await asyncio.to_thread(_square_icon, data, size, THEME_BG.get(c.get("theme", "green"), "#0A0A0A"))
    return Response(content=data, media_type="image/png", headers={"Cache-Control": "public, max-age=3600"})


@api.get("/public/manifest/{slug}")
async def public_manifest(slug: str):
    c = await client_by_slug(slug)
    name = c.get("display_name") or "Marco Polo Wentylacja"
    v = c.get("logo_version")
    icons = ([{"src": f"/api/public/logo/{slug}?size={s}&v={v}", "sizes": f"{s}x{s}", "type": "image/png", "purpose": "any"}
              for s in (192, 512)] if c.get("logo_path") else await site_icons())
    return manifest(name, f"/?klient={slug}", THEME_BG.get(c.get("theme", "green"), "#0A0A0A"), icons)


def manifest(name: str, start: str, bg: str, icons: list) -> JSONResponse:
    return JSONResponse({"id": start, "name": name, "short_name": name, "lang": "pl", "start_url": start, "scope": "/",
                         "display": "standalone", "orientation": "portrait", "background_color": bg, "theme_color": bg,
                         "icons": icons}, media_type="application/manifest+json")


async def site_icons() -> list:
    doc = await db.site_settings.find_one({"_id": "main"}) or {}
    if doc.get("logo_path"):
        return [{"src": f"/api/public/site/logo?size={s}&v={doc['logo_version']}", "sizes": f"{s}x{s}", "type": "image/png",
                 "purpose": "any"} for s in (192, 512)]
    return [{"src": "/icon-192.png", "sizes": "192x192", "type": "image/png"},
            {"src": "/icon-512.png", "sizes": "512x512", "type": "image/png"},
            {"src": "/icon-maskable-512.png", "sizes": "512x512", "type": "image/png", "purpose": "maskable"}]


@api.get("/public/manifest")
async def public_site_manifest():
    return manifest("Marco Polo Wentylacja", "/", "#0A0A0A", await site_icons())


class SiteIn(BaseModel):
    phone: str = Field(default="", max_length=40, pattern=r"^[0-9+ ()-]*$")
    email: str = Field(default="", max_length=254, pattern=r"^([^@\s;,]+@[^@\s;,]+\.[^@\s;,]+)?$")
    address: str = Field(default="", max_length=200)
    website: str = Field(default="", max_length=200, pattern=r"^([A-Za-z0-9.:/_?=&%#-]*)$")


ClientPatch.model_rebuild()


@api.get("/public/site")
async def public_site():
    doc = await db.site_settings.find_one({"_id": "main"}) or {}
    return {"phone": doc.get("phone", ""), "email": doc.get("email", ""),
            "address": doc.get("address", ""), "website": doc.get("website", ""),
            "logo_version": doc.get("logo_version") if doc.get("logo_path") else None}


@api.get("/public/site/logo")
async def public_site_logo(size: int = 0):
    doc = await db.site_settings.find_one({"_id": "main"}) or {}
    if not doc.get("logo_path"):
        raise HTTPException(404, "No logo")
    data = await asyncio.to_thread(get_object, doc["logo_path"])
    if size in (32, 180, 192, 512):
        data = await asyncio.to_thread(_square_icon, data, size, "#0A0A0A")
    return Response(content=data, media_type="image/png", headers={"Cache-Control": "public, max-age=3600"})


@api.post("/site/logo")
async def upload_site_logo(file: UploadFile = File(...), user: dict = Depends(require("superadmin"))):
    path, version = await store_logo(file, "site")
    await db.site_settings.update_one({"_id": "main"}, {"$set": {"logo_path": path, "logo_version": version}}, upsert=True)
    return {"ok": True, "logo_version": version}


@api.delete("/site/logo")
async def delete_site_logo(user: dict = Depends(require("superadmin"))):
    await db.site_settings.update_one({"_id": "main"}, {"$unset": {"logo_path": "", "logo_version": ""}})
    return {"ok": True}


class ContactIn(BaseModel):
    name: str = Field(min_length=2, max_length=80)
    phone: str = Field(default="", max_length=40, pattern=r"^[0-9+ ()-]*$")
    email: str = Field(default="", max_length=254, pattern=r"^([^@\s;,]+@[^@\s;,]+\.[^@\s;,]+)?$")
    message: str = Field(min_length=5, max_length=2000)
    company: str = Field(default="", max_length=200)


@api.post("/public/contact")
async def public_contact(body: ContactIn, request: Request):
    if body.company:
        return {"ok": True}
    body.name = " ".join(body.name.split())
    if len(body.name) < 2:
        raise HTTPException(422, "name too short")
    if not body.phone.strip() and not body.email.strip():
        raise HTTPException(422, "phone or email required")
    site = await db.site_settings.find_one({"_id": "main"}) or {}
    if not site.get("email"):
        raise HTTPException(503, "contact not configured")
    ip = (request.headers.get("x-forwarded-for") or (request.client.host if request.client else "")).split(",")[-1].strip()
    since = datetime.now(timezone.utc) - timedelta(hours=1)
    if (await db.contact_requests.count_documents({"ip": ip, "ts": {"$gte": since}}) >= 3
            or await db.contact_requests.count_documents({"ts": {"$gte": since}}) >= 30):
        raise HTTPException(429, "too many requests")
    await db.contact_requests.insert_one({"ip": ip, "ts": datetime.now(timezone.utc)})
    text = (f"Nowe zapytanie ze strony głównej\n\nImię / firma: {body.name.strip()}\nTelefon: {body.phone.strip() or '-'}\n"
            f"E-mail: {body.email.strip() or '-'}\n\nWiadomość:\n{body.message.strip()}\n")
    st, err = await asyncio.to_thread(notify.deliver, "email", site["email"], f"Zapytanie ze strony – {body.name.strip()[:60]}", text,
                                      (), None, None, body.email.strip() or None)
    await log_event("info" if st == "sent" else "warning", f"Formularz kontaktowy: {body.name.strip()[:60]} – {st}", {"user": "public"})
    if st != "sent":
        raise HTTPException(503, "send failed")
    return {"ok": True}


@api.put("/site")
async def put_site(body: SiteIn, user: dict = Depends(require("superadmin"))):
    await db.site_settings.update_one({"_id": "main"}, {"$set": {"phone": body.phone.strip(), "email": body.email.strip(),
                                                                   "address": body.address.strip(), "website": body.website.strip()}}, upsert=True)
    return await public_site()


@api.get("/branding")
async def branding(client_id: Optional[str] = None, user: dict = Depends(require("viewer"))):
    cid = client_id if user["role"] == "superadmin" else user.get("client_id")
    c = await db.clients.find_one({"_id": ObjectId(cid)}) if cid and ObjectId.is_valid(cid) else None
    if not c:
        return {"client_id": None, "app_name": None, "theme": "green", "logo_version": None}
    return {"client_id": cid, "app_name": c.get("display_name") or None, "theme": c.get("theme", "green"),
            "logo_version": c.get("logo_version") if c.get("logo_path") else None}


@api.delete("/clients/{client_id}")
async def delete_client(client_id: str, user: dict = Depends(require("superadmin"))):
    if not ObjectId.is_valid(client_id):
        raise HTTPException(404, "Client not found")
    if any(d["client_id"] == client_id for d in DEVICES.values()):
        raise HTTPException(409, "Client still has devices")
    res = await db.clients.delete_one({"_id": ObjectId(client_id)})
    if not res.deleted_count:
        raise HTTPException(404, "Client not found")
    await db.users.delete_many({"client_id": client_id})
    await db.alarm_settings.delete_one({"_id": client_id})
    return {"ok": True}


@api.get("/devices")
async def list_devices(user: dict = Depends(require("viewer"))):
    ids = sorted((d for d in DEVICES if visible(user, d)), key=lambda d: (DEVICES[d]["client_name"], DEVICES[d]["name"]))
    return {"devices": [device_view(d, user) for d in ids]}


@api.post("/devices")
async def create_device(body: DeviceIn, user: dict = Depends(require("superadmin"))):
    if not ObjectId.is_valid(body.client_id) or not await db.clients.find_one({"_id": ObjectId(body.client_id)}):
        raise HTTPException(422, "Unknown client")
    if body.device_id in DEVICES:
        raise HTTPException(409, "Device ID already exists")
    password = secrets.token_urlsafe(18)
    await db.devices.insert_one({**body.model_dump(), "name": body.name.strip(), "mqtt_username": body.device_id,
                                 "mqtt_password_enc": encrypt_secret(password), "created_at": now_iso()})
    await load_registry()
    sims.sync()
    dynsec_add_device(body.device_id, password)
    await log_event("info", f"Dodano sterownik: {dev_label(body.device_id)} ({body.chambers} kom.)",
                    {"user": user["email"]}, body.device_id)
    return {**device_view(body.device_id, user), "mqtt_username": body.device_id, "mqtt_password": password}


@api.get("/devices/{device_id}/credentials")
async def device_credentials(device_id: str, user: dict = Depends(require("superadmin"))):
    d = await db.devices.find_one({"device_id": device_id})
    if not d:
        raise HTTPException(404, "Device not found")
    return {"mqtt_username": d.get("mqtt_username", device_id), "mqtt_password": decrypt_secret(d.get("mqtt_password_enc", "")),
            "state_topic": state_topic(device_id), "command_topic": command_topic(device_id),
            "broker_host": MQTT_HOST, "broker_port": MQTT_PORT, "tls": MQTT_TLS}


@api.patch("/devices/{device_id}")
async def patch_device(device_id: str, body: DevicePatch, user: dict = Depends(require("superadmin"))):
    upd = {k: v for k, v in body.model_dump().items() if v is not None}
    res = await db.devices.update_one({"device_id": device_id}, {"$set": upd})
    if not res.matched_count:
        raise HTTPException(404, "Device not found")
    await load_registry()
    sims.sync()
    return device_view(device_id, user)


@api.delete("/devices/{device_id}")
async def delete_device(device_id: str, user: dict = Depends(require("superadmin"))):
    if device_id not in DEVICES:
        raise HTTPException(404, "Device not found")
    label = dev_label(device_id)
    cid = DEVICES.pop(device_id)["client_id"]
    gateway.states.pop(device_id, None)
    gateway.last_seen.pop(device_id, None)
    sims.sync()
    await db.devices.delete_one({"device_id": device_id})
    dynsec_remove_device(device_id)
    for coll in ("device_thresholds", "plc_history", "plc_plant_history", "plc_events"):
        await db[coll].delete_many({"device_id": device_id})
    for k in [k for k in alarm_active if k.startswith(f"{device_id}:")]:
        alarm_active.pop(k)
    await save_alarm_state()
    await load_registry()
    sims.sync()
    gateway.states.pop(device_id, None)
    await log_event("info", f"Usunięto sterownik: {label}", {"user": user["email"]}, client_id=cid)
    return {"ok": True}


class AutoReturnIn(BaseModel):
    time: str = Field(pattern=r"^(([01]\d|2[0-3]):[0-5]\d)?$")


@api.get("/auto-return")
async def get_auto_return(client_id: Optional[str] = None, user: dict = Depends(require("admin"))):
    c = await client_or_404(client_scope(user, client_id))
    return {"time": c.get("auto_return", "00:00")}


@api.put("/auto-return")
async def put_auto_return(body: AutoReturnIn, client_id: Optional[str] = None, user: dict = Depends(require("admin"))):
    c = await client_or_404(client_scope(user, client_id))
    await db.clients.update_one({"_id": c["_id"]}, {"$set": {"auto_return": body.time}})
    await load_registry()
    await log_event("info", f"Powrót do AUTO: {body.time or 'wyłączony'}", {"user": user["email"]}, client_id=str(c["_id"]))
    return {"time": body.time}


@api.post("/admin/auto-reset")
async def admin_auto_reset(user: dict = Depends(require("superadmin"))):
    await reset_manual_to_auto()
    return {"ok": True}


@api.post("/devices/{device_id}/command")
async def device_command(device_id: str, cmd: CommandIn, user: dict = Depends(require("operator"))):
    d = device_or_404(user, device_id)
    if cmd.chamber >= d["chambers"]:
        raise HTTPException(422, "Unknown chamber")
    if cmd.type == "set_setpoint" and (isinstance(cmd.value, bool) or not isinstance(cmd.value, (int, float))):
        raise HTTPException(422, "value must be a number")
    if cmd.type == "set_mode" and cmd.value not in ("auto", "manual"):
        raise HTTPException(422, "value must be auto|manual")
    if cmd.type == "set_vent":
        v = cmd.value
        if not isinstance(v, dict) or not v or any(
                k not in VENT_FIELDS or not _num(x) or not 0 <= x <= 100 for k, x in v.items()):
            raise HTTPException(422, "invalid vent values")
    body = {**cmd.model_dump(), "request_id": uuid.uuid4().hex}
    if not gateway.publish_command(device_id, body):
        raise HTTPException(503, "MQTT publish failed")
    name = f"{d['name']} – {ch_name(device_id, cmd.chamber)}"
    if cmd.type == "set_setpoint":
        msg = f"{name}: nastawa = {cmd.value}°C"
    elif cmd.type == "set_mode":
        msg = f"{name}: tryb {'RĘCZNY' if cmd.value == 'manual' else 'AUTO'}"
    elif cmd.type == "set_vent":
        lbl = {"curtain_left": "kurtyna L zamkn.", "curtain_right": "kurtyna P zamkn.", "chimney": "komin otw."}
        msg = f"{name}: " + ", ".join(
            f"{lbl[k]} {round(100 - x) if k != 'chimney' else round(x)}%" for k, x in cmd.value.items())
    else:
        msg = f"{name}: {'włączona' if cmd.value else 'wyłączona'}"
    await log_event("command", msg, {**body, "user": user["email"]}, device_id)
    return {"queued": True, "request_id": body["request_id"]}


@api.get("/devices/{device_id}/thresholds")
async def get_thresholds(device_id: str, user: dict = Depends(require("viewer"))):
    d = device_or_404(user, device_id)
    return {"thresholds": [{"chamber": i, **th_for(device_id, i)} for i in range(d["chambers"])]}


@api.put("/devices/{device_id}/thresholds/{chamber}")
async def set_threshold(device_id: str, chamber: int, body: ThresholdIn, user: dict = Depends(require("operator"))):
    d = device_or_404(user, device_id)
    if not 0 <= chamber < d["chambers"]:
        raise HTTPException(404, "unknown chamber")
    th = {"temp_max": round(body.temp_max, 1), "co2_max": float(round(body.co2_max))}
    THRESHOLDS[(device_id, chamber)] = th
    await db.device_thresholds.update_one({"device_id": device_id, "chamber": chamber}, {"$set": th}, upsert=True)
    await log_event("info", f"{d['name']} – {ch_name(device_id, chamber)}: progi temp ≤ {th['temp_max']}°C, CO₂ ≤ {int(th['co2_max'])} ppm",
                    {"user": user["email"]}, device_id)
    st = gateway.states.get(device_id)
    if st:
        prev = copy.deepcopy(st)
        annotate(device_id, st)
        gateway.detect_alarms(device_id, prev, st)
    return {"ok": True, "thresholds": th}


class ChamberNameIn(BaseModel):
    name: str = Field(default="", max_length=30)


@api.put("/devices/{device_id}/chambers/{chamber}/name")
async def rename_chamber(device_id: str, chamber: int, body: ChamberNameIn, user: dict = Depends(require("admin"))):
    d = device_or_404(user, device_id)
    if not 0 <= chamber < d["chambers"]:
        raise HTTPException(404, "unknown chamber")
    name, old = " ".join(body.name.split()), ch_name(device_id, chamber)
    if (name or f"Komora {chamber + 1}") == old:
        return {"ok": True, "name": old}
    op = {"$set": {f"chamber_names.{chamber}": name}} if name else {"$unset": {f"chamber_names.{chamber}": ""}}
    await db.devices.update_one({"device_id": device_id}, op)
    await load_registry()
    st = gateway.states.get(device_id)
    if st:
        st["chambers"][chamber]["name"] = ch_name(device_id, chamber)
    await log_event("info", f"{d['name']}: zmieniono nazwę „{old}” → „{ch_name(device_id, chamber)}”", {"user": user["email"]}, device_id)
    return {"ok": True, "name": ch_name(device_id, chamber)}


@api.get("/devices/{device_id}/history")
async def history(device_id: str, hours: int = 3, user: dict = Depends(require("viewer"))):
    device_or_404(user, device_id)
    since = (datetime.now(timezone.utc) - timedelta(hours=max(1, min(24, hours)))).isoformat()
    q = {"device_id": device_id, "ts": {"$gte": since}}
    docs = await db.plc_history.find(q, {"_id": 0}).sort("ts", 1).to_list(20000)
    out: dict = {}
    for doc in docs:
        out.setdefault(str(doc["chamber"]), []).append(doc)
    plant = await db.plc_plant_history.find(q, {"_id": 0}).sort("ts", 1).to_list(3000)
    return {"history": out, "plant": plant}


@api.get("/devices/{device_id}/usage/daily")
async def usage_daily(device_id: str, days: int = 7, user: dict = Depends(require("viewer"))):
    device_or_404(user, device_id)
    days = max(1, min(60, days))
    start = (datetime.now(timezone.utc) - timedelta(days=days - 1)).date()
    rows = await db.plc_plant_history.aggregate([
        {"$match": {"device_id": device_id, "ts": {"$gte": start.isoformat()}}},
        {"$group": {"_id": {"$substr": ["$ts", 0, 10]}, "p": {"$sum": "$power_kw"}, "w": {"$sum": "$water_lpm"}}},
    ]).to_list(100)
    by = {r["_id"]: r for r in rows}
    meters = {m["day"]: m async for m in db.meter_daily.find({"device_id": device_id, "day": {"$gte": start.isoformat()}})}
    out = []
    for i in range(days):
        day = (start + timedelta(days=i)).isoformat()
        r = by.get(day, {"p": 0, "w": 0})
        m = meters.get(day)
        out.append({"date": day, "kwh": round(m.get("kwh", 0), 2) if m else round(r["p"] * SAMPLE_S / 3600, 2),
                    "water_m3": round(sum(m.get("water", {}).values()) / 1000, 3) if m else round(r["w"] * SAMPLE_S / 60 / 1000, 3)})
    return {"days": out}


DATE_RX = r"^\d{4}-\d{2}-\d{2}$"


def parse_day(v: str) -> str:
    try:
        return datetime.strptime(v, "%Y-%m-%d").date().isoformat()
    except (ValueError, TypeError):
        raise HTTPException(422, "date must be YYYY-MM-DD")


def range_or_422(start: str, end: str) -> tuple:
    a, b = parse_day(start), parse_day(end)
    if a > b or (datetime.fromisoformat(b) - datetime.fromisoformat(a)).days > 1100:
        raise HTTPException(422, "invalid date range")
    return a, b


@api.get("/devices/{device_id}/usage/range")
async def usage_range(device_id: str, start: str, end: str, user: dict = Depends(require("viewer"))):
    device_or_404(user, device_id)
    a, b = range_or_422(start, end)
    return await usage_between(device_id, a, b)


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


DEFAULT_REPORT_LOGO = (ROOT_DIR / "assets" / "marco-polo.png").read_bytes()
LOGO_ROWS = 5


def report_logo_png(raw: bytes, height: int) -> bytes:
    im = Image.open(io.BytesIO(raw)).convert("RGBA")
    im = im.resize((max(1, round(im.width * height / im.height)), height), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


async def client_logo(client_id: Optional[str]) -> bytes:
    c = await db.clients.find_one({"_id": ObjectId(client_id)}) if client_id and ObjectId.is_valid(client_id) else None
    if c and c.get("logo_path"):
        try:
            return await asyncio.to_thread(get_object, c["logo_path"])
        except Exception as e:
            logger.warning("client logo unavailable: %s", e)
    return DEFAULT_REPORT_LOGO


FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def csv_cell(v) -> str:
    if v is None:
        return ""
    if _num(v):
        return str(v).replace(".", ",")
    v = str(v)
    return ("'" + v if v[:1] in FORMULA_START else v).replace(";", " ").replace("\n", " ").replace("\r", " ")


def xlsx_bytes(sheets: list, logo: Optional[bytes] = None) -> bytes:
    """sheets: [(title, rows, col_widths)] – first data row of each sheet is bold header; logo goes above it."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.drawing.image import Image as XlImage
    png = report_logo_png(logo, 80) if logo else None
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows, widths in sheets:
        ws = wb.create_sheet(title)
        if png:
            for _ in range(LOGO_ROWS):
                ws.append([])
            ws.add_image(XlImage(io.BytesIO(png)), "A1")
        for r in rows:
            ws.append(r)
            for c in ws[ws.max_row]:
                if isinstance(c.value, str) and c.value[:1] in FORMULA_START:
                    c.data_type = "s"
        for c in ws[LOGO_ROWS + 1 if png else 1] if rows else []:
            c.font = Font(bold=True)
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def usage_xlsx(device_id: str, a: str, b: str) -> bytes:
    d = DEVICES[device_id]
    u = await usage_between(device_id, a, b)
    names = [ch_name(device_id, i) for i in range(d["chambers"])]
    summary = [[f"{d['client_name']} – {d['name']}", f"{a} – {b}", ""], ["Pozycja", "Zużycie", "Jednostka"],
               ["Prąd (cała chlewnia)", u["kwh"], "kWh"], ["Woda razem", round(u["water_total_l"] / 1000, 3), "m³"]]
    summary += [[f"Woda – {nm}", round(w / 1000, 3), "m³"] for nm, w in zip(names, u["water_l"])]
    r = await client_rates(d["client_id"])
    if r.get("kwh") is not None or r.get("m3") is not None:
        summary += [[], ["Koszty", "Kwota", "Stawka"]]
        if r.get("kwh") is not None:
            summary.append(["Koszt prądu", money(u["kwh"] * r["kwh"]), f"{r['kwh']} zł/kWh"])
        if r.get("m3") is not None:
            summary.append(["Koszt wody razem", money(u["water_total_l"] / 1000 * r["m3"]), f"{r['m3']} zł/m³"])
            summary += [[f"Koszt wody – {nm}", money(w / 1000 * r["m3"]), "zł"] for nm, w in zip(names, u["water_l"])]
        summary.append(["KOSZT RAZEM", money(u["kwh"] * (r.get("kwh") or 0) + u["water_total_l"] / 1000 * (r.get("m3") or 0)), "zł"])
    days = [["Dzień", "Prąd kWh"] + [f"Woda {nm} m³" for nm in names]]
    days += [[r["day"], r["kwh"]] + [round(w / 1000, 3) for w in r["water_l"]] for r in u["days"]]
    logo = await client_logo(d["client_id"])
    return await asyncio.to_thread(xlsx_bytes, [("Podsumowanie", summary, {"A": 34, "B": 24}), ("Dni", days, {"A": 14})], logo)


def file_response(data: bytes, fname: str, mime: str) -> Response:
    return Response(content=data, media_type=mime, headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@api.get("/devices/{device_id}/usage/range/xlsx")
async def usage_range_xlsx(device_id: str, start: str, end: str, user: dict = Depends(require("viewer"))):
    device_or_404(user, device_id)
    a, b = range_or_422(start, end)
    return file_response(await usage_xlsx(device_id, a, b), f"zuzycie_{device_id}_{a}_{b}.xlsx", XLSX_MIME)


@api.get("/devices/{device_id}/usage/range/csv")
async def usage_range_csv(device_id: str, start: str, end: str, user: dict = Depends(require("viewer"))):
    d = device_or_404(user, device_id)
    a, b = range_or_422(start, end)
    u = await usage_between(device_id, a, b)
    names = [csv_cell(ch_name(device_id, i)).lstrip("'") for i in range(d["chambers"])]
    lines = [";".join(["dzien", "prad_kWh"] + [f"woda_{nm}_m3" for nm in names])]
    lines += [";".join([r["day"], str(r["kwh"]).replace(".", ",")] + [str(round(w / 1000, 3)).replace(".", ",") for w in r["water_l"]])
              for r in u["days"]]
    lines.append(";".join(["RAZEM", str(u["kwh"]).replace(".", ",")] + [str(round(w / 1000, 3)).replace(".", ",") for w in u["water_l"]]))
    return file_response(("\ufeff" + "\n".join(lines) + "\n").encode(), f"zuzycie_{device_id}_{a}_{b}.csv", "text/csv; charset=utf-8")


def prev_month(today: Optional[datetime] = None) -> tuple:
    first = (today or datetime.now(PL_TZ)).date().replace(day=1)
    last = first - timedelta(days=1)
    return last.replace(day=1).isoformat(), last.isoformat()


async def send_monthly_report(client_id: str, force: bool = False) -> dict:
    a, b = prev_month()
    key = f"{client_id}|{a[:7]}"
    if not force and await db.monthly_reports.find_one({"_id": key}):
        return {"skipped": True}
    devs = [k for k, v in DEVICES.items() if v["client_id"] == client_id]
    to = [u["email"] async for u in db.users.find({"client_id": client_id, "role": "admin"})]
    if not devs or not to:
        return {"sent": 0, "recipients": to}
    files = tuple([(f"zuzycie_{dv}_{a[:7]}.xlsx", await usage_xlsx(dv, a, b)) for dv in devs])
    name = DEVICES[devs[0]]["client_name"]
    r = await client_rates(client_id)
    cost = ""
    if r.get("kwh") is not None or r.get("m3") is not None:
        us = [await usage_between(dv, a, b) for dv in devs]
        cp = sum(u["kwh"] for u in us) * (r.get("kwh") or 0)
        cw = sum(u["water_total_l"] for u in us) / 1000 * (r.get("m3") or 0)
        cost = f"Koszt prądu: {cp:.2f} zł, koszt wody: {cw:.2f} zł, razem: {cp + cw:.2f} zł."
    body = (f"Dzień dobry,\n\nw załączniku raport zużycia wody i prądu za okres {a} – {b} ({name}).\n"
            + (f"{cost}\n" if cost else "") + "\nMarco Polo Wentylacja")
    from html import escape
    html = (f'<div style="font-family:Arial,sans-serif;color:#222"><img src="cid:logo" alt="" style="max-height:80px"><br><br>'
            f'Dzień dobry,<br><br>w załączniku raport zużycia wody i prądu za okres <b>{a} – {b}</b> ({escape(name)}).'
            + (f'<br><b>{cost}</b>' if cost else '') + '<br><br>Marco Polo Wentylacja</div>')
    inline = report_logo_png(await client_logo(client_id), 120)
    res = []
    for r in to:
        res.append(await asyncio.to_thread(notify.deliver, "email", r, f"Raport miesięczny {a[:7]} – {name}", body, files,
                                           html, inline))
    await db.monthly_reports.update_one({"_id": key}, {"$set": {"ts": now_iso(), "to": to, "status": [x[0] for x in res]}}, upsert=True)
    await log_event("info", f"Raport miesięczny {a[:7]}: " + ", ".join(f"{e} – {st}" for e, (st, _) in zip(to, res)), {"user": "system"},
                    client_id=client_id)
    return {"sent": sum(x[0] == "sent" for x in res), "recipients": to, "status": [x[0] for x in res]}


async def monthly_tick():
    now = datetime.now(PL_TZ)
    if now.day != 1 or now.strftime("%H:%M") != "06:00":
        return
    async for c in db.clients.find({"monthly_report": {"$ne": False}}):
        try:
            await send_monthly_report(str(c["_id"]))
        except Exception as e:
            logger.error("monthly report failed: %s", e)


@api.post("/monthly-report/send")
async def monthly_report_now(client_id: Optional[str] = None, user: dict = Depends(require("admin"))):
    c = await client_or_404(client_scope(user, client_id))
    return await send_monthly_report(str(c["_id"]), force=True)


@api.get("/monthly-report")
async def monthly_report_get(client_id: Optional[str] = None, user: dict = Depends(require("admin"))):
    c = await client_or_404(client_scope(user, client_id))
    to = [u["email"] async for u in db.users.find({"client_id": str(c["_id"]), "role": "admin"})]
    return {"enabled": c.get("monthly_report", True), "recipients": to, "period": prev_month()}


class MonthlyIn(BaseModel):
    enabled: bool


@api.put("/monthly-report")
async def monthly_report_put(body: MonthlyIn, client_id: Optional[str] = None, user: dict = Depends(require("admin"))):
    c = await client_or_404(client_scope(user, client_id))
    await db.clients.update_one({"_id": c["_id"]}, {"$set": {"monthly_report": body.enabled}})
    return {"enabled": body.enabled}


@api.get("/client-contact")
async def get_client_contact(client_id: Optional[str] = None, user: dict = Depends(require("admin"))):
    return client_contact(await client_or_404(client_scope(user, client_id)))


@api.put("/client-contact")
async def put_client_contact(body: SiteIn, client_id: Optional[str] = None, user: dict = Depends(require("admin"))):
    c = await client_or_404(client_scope(user, client_id))
    contact = {k: v.strip() for k, v in body.model_dump().items()}
    await db.clients.update_one({"_id": c["_id"]}, {"$set": {"contact": contact}})
    await log_event("info", "Zmieniono dane kontaktowe klienta", {"user": user["email"]}, client_id=str(c["_id"]))
    return contact


class RatesIn(BaseModel):
    kwh: Optional[float] = Field(default=None, ge=0, le=100)
    m3: Optional[float] = Field(default=None, ge=0, le=1000)


@api.get("/rates")
async def get_rates(client_id: Optional[str] = None, user: dict = Depends(require("admin"))):
    c = await client_or_404(client_scope(user, client_id))
    r = c.get("rates") or {}
    return {"kwh": r.get("kwh"), "m3": r.get("m3")}


@api.put("/rates")
async def put_rates(body: RatesIn, client_id: Optional[str] = None, user: dict = Depends(require("admin"))):
    c = await client_or_404(client_scope(user, client_id))
    await db.clients.update_one({"_id": c["_id"]}, {"$set": {"rates": {"kwh": body.kwh, "m3": body.m3}}})
    await log_event("info", f"Stawki: prąd {body.kwh if body.kwh is not None else '-'} zł/kWh, woda {body.m3 if body.m3 is not None else '-'} zł/m³",
                    {"user": user["email"]}, client_id=str(c["_id"]))
    return {"kwh": body.kwh, "m3": body.m3}


class CycleIn(BaseModel):
    start: str = Field(pattern=DATE_RX)
    animals: Optional[int] = Field(default=None, ge=0, le=100000)


class CycleEndIn(BaseModel):
    end: Optional[str] = Field(default=None, pattern=DATE_RX)


async def client_rates(client_id: Optional[str]) -> dict:
    c = await db.clients.find_one({"_id": ObjectId(client_id)}, {"rates": 1}) if client_id and ObjectId.is_valid(client_id) else None
    return (c or {}).get("rates") or {}


async def cycle_ctx(device_id: str) -> dict:
    cycles = await db.cycles.find({"device_id": device_id}, {"chamber": 1, "start": 1, "end": 1}).to_list(5000)
    days = {r["day"]: r.get("kwh", 0) async for r in db.meter_daily.find({"device_id": device_id}, {"day": 1, "kwh": 1})}
    return {"cycles": cycles, "days": days, "rates": await client_rates(DEVICES[device_id]["client_id"])}


def chamber_kwh(c: dict, end: str, ctx: dict) -> float:
    """Plant kWh of each day split equally between chambers with a running cycle that day."""
    today, total = local_day(), 0.0
    for day, kwh in ctx["days"].items():
        if not (c["start"] <= day <= end) or not kwh:
            continue
        n = len({x["chamber"] for x in ctx["cycles"] if x["start"] <= day <= (x.get("end") or today)})
        total += kwh / max(n, 1)
    return round(total, 2)


def money(v: Optional[float]) -> Optional[float]:
    return None if v is None else round(v, 2)


async def cycle_view(c: dict, ctx: Optional[dict] = None) -> dict:
    ctx = ctx or await cycle_ctx(c["device_id"])
    end = c.get("end") or local_day()
    u = await usage_between(c["device_id"], c["start"], end) if c["start"] <= end else {"kwh": 0, "water_l": []}
    w = u["water_l"][c["chamber"]] if c["chamber"] < len(u["water_l"]) else 0
    kwh_ch = chamber_kwh(c, end, ctx)
    r = ctx["rates"]
    cw = w / 1000 * r["m3"] if r.get("m3") is not None else None
    cp = kwh_ch * r["kwh"] if r.get("kwh") is not None else None
    tot = None if cw is None and cp is None else (cw or 0) + (cp or 0)
    return {"id": str(c["_id"]), "chamber": c["chamber"], "start": c["start"], "end": c.get("end"), "animals": c.get("animals"),
            "days": (datetime.fromisoformat(end) - datetime.fromisoformat(c["start"])).days + 1,
            "water_l": w, "water_per_head_l": round(w / c["animals"], 2) if c.get("animals") else None,
            "kwh": u["kwh"], "kwh_chamber": kwh_ch, "cost_water": money(cw), "cost_power": money(cp), "cost_total": money(tot),
            "cost_per_head": round(tot / c["animals"], 3) if tot is not None and c.get("animals") else None,
            "started_by": c.get("started_by"), "ended_by": c.get("ended_by")}


@api.get("/devices/{device_id}/cycles")
async def list_cycles(device_id: str, chamber: int, user: dict = Depends(require("viewer"))):
    device_or_404(user, device_id)
    rows = await db.cycles.find({"device_id": device_id, "chamber": chamber}).sort("start", -1).to_list(100)
    ctx = await cycle_ctx(device_id)
    return {"cycles": [await cycle_view(c, ctx) for c in rows]}


CYCLE_HEAD = ["Komora", "Wstawienie", "Koniec", "Dni", "Sztuk", "Woda m³", "Woda na sztukę l", "l/szt./dzień",
              "Prąd kWh (cała chlewnia)", "Prąd kWh (komora)", "Koszt wody zł", "Koszt prądu zł", "Koszt razem zł",
              "Koszt na sztukę zł", "Wpisał", "Zakończył"]


async def cycle_rows(device_id: str, chamber: Optional[int]) -> list:
    q = {"device_id": device_id, **({"chamber": chamber} if chamber is not None else {})}
    out, ctx = [], await cycle_ctx(device_id)
    for c in await db.cycles.find(q).sort([("chamber", 1), ("start", -1)]).to_list(2000):
        v = await cycle_view(c, ctx)
        ph = v["water_per_head_l"]
        out.append([ch_name(device_id, v["chamber"]), v["start"], v["end"] or "trwa", v["days"], v["animals"],
                    round(v["water_l"] / 1000, 3), ph, round(ph / v["days"], 3) if ph is not None else None, v["kwh"],
                    v["kwh_chamber"], v["cost_water"], v["cost_power"], v["cost_total"], v["cost_per_head"],
                    v["started_by"], v["ended_by"]])
    return out


def avg(vals: list) -> Optional[float]:
    vals = [v for v in vals if v is not None]
    return sum(vals) / len(vals) if vals else None


async def chamber_compare(device_id: str) -> list:
    d, ctx = DEVICES[device_id], await cycle_ctx(device_id)
    done = await db.cycles.find({"device_id": device_id, "end": {"$ne": None}}).to_list(5000)
    views = [await cycle_view(c, ctx) for c in done]
    out = []
    for ch in range(d["chambers"]):
        vs = [v for v in views if v["chamber"] == ch]
        cph, wph = avg([v["cost_per_head"] for v in vs]), avg([v["water_per_head_l"] for v in vs])
        out.append({"chamber": ch, "name": ch_name(device_id, ch), "cycles": len(vs),
                    "cost_per_head": round(cph, 3) if cph is not None else None,
                    "water_per_head_l": round(wph, 2) if wph is not None else None,
                    "days": round(avg([v["days"] for v in vs]), 1) if vs else None})
    return out


@api.get("/devices/{device_id}/cycles/compare")
async def cycles_compare(device_id: str, user: dict = Depends(require("viewer"))):
    device_or_404(user, device_id)
    return {"chambers": await chamber_compare(device_id)}


async def farm_compare() -> list:
    out = []
    async for c in db.clients.find({}).sort("name", 1):
        cid = str(c["_id"])
        devs = [k for k, d in DEVICES.items() if d["client_id"] == cid]
        vs = []
        for dv in devs:
            ctx = await cycle_ctx(dv)
            vs += [await cycle_view(x, ctx) for x in await db.cycles.find({"device_id": dv, "end": {"$ne": None}}).to_list(5000)]
        r3 = lambda v, n=3: round(v, n) if v is not None else None
        out.append({"client_id": cid, "name": c.get("display_name") or c["name"], "devices": len(devs), "cycles": len(vs),
                    "cost_per_head": r3(avg([v["cost_per_head"] for v in vs])),
                    "water_per_head_l": r3(avg([v["water_per_head_l"] for v in vs]), 2),
                    "kwh_per_head": r3(avg([v["kwh_chamber"] / v["animals"] if v.get("animals") else None for v in vs]))})
    return out


@api.get("/devices/{device_id}/mqtt-debug")
async def mqtt_debug(device_id: str, user: dict = Depends(require("superadmin"))):
    device_or_404(user, device_id)
    return {"online": gateway.is_online(device_id), "broker_connected": gateway.connected,
            "state_topic": state_topic(device_id), "command_topic": command_topic(device_id),
            "last_message": gateway.raw.get(device_id), "last_command": gateway.last_cmd.get(device_id)}


@api.get("/farms/compare")
async def farms_compare(user: dict = Depends(require("superadmin"))):
    return {"farms": await farm_compare()}


@api.get("/farms/compare/xlsx")
async def farms_compare_xlsx(user: dict = Depends(require("superadmin"))):
    rows = [["Ferma", "Sterowniki", "Zakończone rzuty", "Śr. koszt na sztukę zł", "Śr. woda na sztukę l", "Śr. prąd na sztukę kWh"]]
    rows += [[f["name"], f["devices"], f["cycles"], f["cost_per_head"], f["water_per_head_l"], f["kwh_per_head"]] for f in await farm_compare()]
    data = await asyncio.to_thread(xlsx_bytes, [("Porównanie ferm", rows, {"A": 28, "B": 11, "C": 17, "D": 22, "E": 20, "F": 22})], DEFAULT_REPORT_LOGO)
    return file_response(data, f"porownanie_ferm_{local_day()}.xlsx", XLSX_MIME)


@api.get("/devices/{device_id}/cycles/export")
async def cycles_export(device_id: str, fmt: str = "xlsx", chamber: Optional[int] = None, user: dict = Depends(require("viewer"))):
    d = device_or_404(user, device_id)
    rows = await cycle_rows(device_id, chamber)
    fname = f"chowy_{device_id}{'' if chamber is None else f'_komora{chamber + 1}'}_{local_day()}"
    if fmt == "csv":
        txt = "\n".join(";".join(csv_cell(x) for x in r) for r in [CYCLE_HEAD] + rows)
        return file_response(("\ufeff" + txt + "\n").encode(), fname + ".csv", "text/csv; charset=utf-8")
    widths = {"A": 18, "B": 12, "C": 12, "D": 6, "E": 8, "F": 10, "G": 16, "H": 12, "I": 22, "J": 16, "K": 14, "L": 14,
              "M": 15, "N": 17, "O": 26, "P": 26}
    cmp = [["Komora", "Zakończone rzuty", "Śr. koszt na sztukę zł", "Śr. woda na sztukę l", "Śr. długość rzutu dni"]]
    cmp += [[x["name"], x["cycles"], x["cost_per_head"], x["water_per_head_l"], x["days"]] for x in await chamber_compare(device_id)]
    data = await asyncio.to_thread(xlsx_bytes, [("Chowy", [CYCLE_HEAD] + rows, widths),
                                                ("Porównanie komór", cmp, {"A": 20, "B": 17, "C": 22, "D": 20, "E": 21})],
                                   await client_logo(d["client_id"]))
    return file_response(data, fname + ".xlsx", XLSX_MIME)


@api.post("/devices/{device_id}/chambers/{chamber}/cycles")
async def start_cycle(device_id: str, chamber: int, body: CycleIn, user: dict = Depends(require("operator"))):
    d = device_or_404(user, device_id)
    if not 0 <= chamber < d["chambers"]:
        raise HTTPException(404, "unknown chamber")
    start = parse_day(body.start)
    if start > local_day():
        raise HTTPException(422, "start date in the future")
    if await db.cycles.find_one({"device_id": device_id, "chamber": chamber, "end": None}):
        raise HTTPException(409, "active cycle exists")
    doc = {"device_id": device_id, "chamber": chamber, "start": start, "end": None, "active": True, "animals": body.animals,
           "started_by": user["email"], "created_at": now_iso()}
    try:
        await db.cycles.insert_one(doc)
    except DuplicateKeyError:
        raise HTTPException(409, "active cycle exists")
    await log_event("info", f"{d['name']} – {ch_name(device_id, chamber)}: nowe wstawienie {start}"
                    + (f", {body.animals} szt." if body.animals is not None else "") + f" ({user['email']})",
                    {"user": user["email"]}, device_id)
    return await cycle_view(doc)


@api.post("/devices/{device_id}/cycles/{cycle_id}/end")
async def end_cycle(device_id: str, cycle_id: str, body: CycleEndIn, user: dict = Depends(require("operator"))):
    d = device_or_404(user, device_id)
    if not ObjectId.is_valid(cycle_id):
        raise HTTPException(404, "cycle not found")
    c = await db.cycles.find_one({"_id": ObjectId(cycle_id), "device_id": device_id})
    if not c:
        raise HTTPException(404, "cycle not found")
    if c.get("end"):
        raise HTTPException(409, "cycle already ended")
    end = parse_day(body.end) if body.end else local_day()
    if end < c["start"] or end > local_day():
        raise HTTPException(422, "invalid end date")
    await db.cycles.update_one({"_id": c["_id"]}, {"$set": {"end": end, "ended_by": user["email"]}, "$unset": {"active": ""}})
    c.update(end=end, ended_by=user["email"])
    v = await cycle_view(c)
    await log_event("info", f"{d['name']} – {ch_name(device_id, c['chamber'])}: koniec chowu {c['start']} – {end}, "
                    f"woda {round(v['water_l'] / 1000, 2)} m³ ({user['email']})", {"user": user["email"]}, device_id)
    return v


@api.get("/devices/{device_id}/history/xlsx")
async def history_xlsx(device_id: str, chamber: int, hours: int = 24, user: dict = Depends(require("viewer"))):
    d = device_or_404(user, device_id)
    if not 0 <= chamber < d["chambers"]:
        raise HTTPException(404, "unknown chamber")
    hours = max(1, min(24 * 31, hours))
    since = (datetime.now(timezone.utc) - timedelta(hours=hours)).isoformat()
    rows = [["Czas", "Temperatura °C", "Wilgotność %", "CO₂ ppm", "Kurtyna L zamkn. %", "Kurtyna P zamkn. %", "Komin otw. %"]]
    async for doc in db.plc_history.find({"device_id": device_id, "chamber": chamber, "ts": {"$gte": since}}, {"_id": 0}).sort("ts", 1):
        cl, cr = doc.get("curtain_left"), doc.get("curtain_right")
        rows.append([datetime.fromisoformat(doc["ts"]).astimezone(PL_TZ).strftime("%Y-%m-%d %H:%M:%S"), doc.get("temp"),
                     doc.get("humidity"), doc.get("co2"), None if cl is None else round(100 - cl, 1),
                     None if cr is None else round(100 - cr, 1), doc.get("chimney")])
    title = ch_name(device_id, chamber)[:31].translate(str.maketrans("[]:*?/\\", "       "))
    fname = f"{device_id}_komora{chamber + 1}_{hours}h_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return file_response(await asyncio.to_thread(xlsx_bytes, [(title, rows, {"A": 20, "B": 15, "C": 14, "D": 10, "E": 18, "F": 18, "G": 13})],
                                                 await client_logo(d["client_id"])),
                         fname, XLSX_MIME)


@api.get("/devices/{device_id}/history/csv")
async def history_csv(device_id: str, chamber: int, hours: int = 24, user: dict = Depends(require("viewer"))):
    d = device_or_404(user, device_id)
    if not 0 <= chamber < d["chambers"]:
        raise HTTPException(404, "unknown chamber")
    hours = max(1, min(24 * 31, hours))
    since = (datetime.now(timezone.utc) - timedelta(hours=hours)).isoformat()
    cursor = db.plc_history.find({"device_id": device_id, "chamber": chamber, "ts": {"$gte": since}}, {"_id": 0}).sort("ts", 1)

    fmt = csv_cell

    async def gen():
        yield "\ufeffczas;temperatura_C;wilgotnosc_%;co2_ppm;kurtyna_L_zamkn_%;kurtyna_P_zamkn_%;komin_otw_%\n"
        async for doc in cursor:
            cl, cr = doc.get("curtain_left"), doc.get("curtain_right")
            ts = datetime.fromisoformat(doc["ts"]).astimezone(ZoneInfo("Europe/Warsaw")).strftime("%Y-%m-%d %H:%M:%S")
            yield ";".join([ts, fmt(doc.get("temp")), fmt(doc.get("humidity")), fmt(doc.get("co2")),
                            fmt(None if cl is None else round(100 - cl, 1)),
                            fmt(None if cr is None else round(100 - cr, 1)), fmt(doc.get("chimney"))]) + "\n"

    fname = f"{device_id}_komora{chamber + 1}_{hours}h_{datetime.now().strftime('%Y%m%d_%H%M')}.csv"
    return StreamingResponse(gen(), media_type="text/csv; charset=utf-8",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@api.get("/events")
async def events(device_id: Optional[str] = None, limit: int = 100, user: dict = Depends(require("viewer"))):
    q: dict = {} if user["role"] == "superadmin" else {"client_id": user.get("client_id")}
    if device_id:
        d = device_or_404(user, device_id)
        q = {"client_id": d["client_id"], "$or": [{"device_id": device_id}, {"device_id": None}]}
    docs = await db.plc_events.find(q, {"_id": 0}).sort("ts", -1).to_list(max(1, min(500, limit)))
    return {"events": docs}


@api.get("/alarm-settings")
async def get_alarm_settings(client_id: Optional[str] = None, user: dict = Depends(require("admin"))):
    cid = client_scope(user, client_id)
    cfg = await db.alarm_settings.find_one({"_id": cid}) or {}
    sent = await db.alarm_notifications.find({"client_id": cid}, {"_id": 0}).sort("ts", -1).to_list(10)
    return {"client_id": cid, "enabled": cfg.get("enabled", True), "recipients": cfg.get("recipients", []),
            "email_configured": notify.email_configured(), "sms_configured": notify.sms_configured(), "reminder_min": ALARM_REMIND_S // 60, "recent": sent}


@api.put("/alarm-settings")
async def put_alarm_settings(body: AlarmSettingsIn, client_id: Optional[str] = None, user: dict = Depends(require("admin"))):
    cid = client_scope(user, client_id)
    recs = []
    for r in body.recipients:
        v = r.value.strip().lower() if r.type == "email" else re.sub(r"[\s-]", "", r.value)
        if not re.match(EMAIL_RX if r.type == "email" else PHONE_RX, v):
            raise HTTPException(422, f"invalid {r.type}: {r.value}")
        if not any(x["type"] == r.type and x["value"] == v for x in recs):
            recs.append({"type": r.type, "value": v})
    await db.alarm_settings.update_one({"_id": cid}, {"$set": {"enabled": body.enabled, "recipients": recs}}, upsert=True)
    await log_event("info", f"Powiadomienia o alarmach {'WŁ.' if body.enabled else 'WYŁ.'}, odbiorcy: {len(recs)}",
                    {"user": user["email"]}, client_id=cid)
    return await get_alarm_settings(cid, user)


@api.post("/alarm-settings/test")
async def test_alarm(client_id: Optional[str] = None, user: dict = Depends(require("admin"))):
    cid = client_scope(user, client_id)
    cfg = await db.alarm_settings.find_one({"_id": cid}) or {}
    if not cfg.get("enabled", True) or not cfg.get("recipients"):
        raise HTTPException(409, "Notifications disabled or no recipients")
    await send_alarm(cid, "TEST: Marco Polo Wentylacja", f"Powiadomienie testowe wysłane przez {user['email']}", "test")
    return await get_alarm_settings(cid, user)


def origin_allowed(origin: Optional[str], host: Optional[str]) -> bool:
    if not origin:
        return True
    origin = origin.rstrip("/")
    return origin == FRONTEND_URL or (bool(host) and origin.split("://", 1)[-1] == host)


@app.websocket("/api/ws")
async def websocket_endpoint(ws: WebSocket):
    user = await auth.user_from_token(ws.cookies.get("access_token"))
    if not user or not origin_allowed(ws.headers.get("origin"), ws.headers.get("host")):
        await ws.close(code=4401)
        return
    await ws.accept()
    gateway.ws_clients[ws] = user
    for device_id, st in list(gateway.states.items()):
        if visible(user, device_id):
            await ws.send_json({"type": "state", "payload": st})
    try:
        while True:
            await ws.receive_text()
    except Exception:
        pass
    finally:
        gateway.ws_clients.pop(ws, None)


@app.middleware("http")
async def origin_guard(request: Request, call_next):
    if request.method in ("POST", "PUT", "PATCH", "DELETE"):
        if not origin_allowed(request.headers.get("origin"), request.headers.get("host")):
            return JSONResponse({"detail": "Forbidden origin"}, status_code=403)
    return await call_next(request)


app.include_router(auth.router)
app.include_router(api)
app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=[FRONTEND_URL],
    allow_methods=["*"],
    allow_headers=["*"],
)
